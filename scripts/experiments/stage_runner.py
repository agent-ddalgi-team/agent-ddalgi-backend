"""Offline saved-draft stages: render, page images, preservation and regression.

Run with the repository's Python, from the product revision being compared::

    python -B scripts/experiments/stage_runner.py stage2 --manifest INPUT.json \
        --run-dir OUTPUT --expected-template template_v11

The private manifest has ``documents`` entries with unique ``label``,
``document_path`` and ``assets_dir``. Paths may be absolute or relative to the
manifest. Asset directories contain asset_manifest.json and archived images;
stage1 freezes unmodified drafts in baseline/. stage2 uses the current product's
pagination only; --previous-dir and --baseline-dir can select explicit inputs.
Stage3/4 caption/body transformations and caption font trials are deferred until
team policy decisions. --docx-label and --smoke-four select additional checks.
Target page differences are diagnostics, not a new product approval policy.

No API setup, live database, implicit retries, Git writes or private finalizer.
The runner never deletes existing output. An incomplete document directory must
be archived explicitly before --resume. Resume requires identical code, tests,
manifest, input and previously completed artifact hashes. Regression is always
run afresh; historical logs cannot be reused by this runner.
"""

from __future__ import annotations

import argparse
from contextlib import contextmanager, redirect_stderr, redirect_stdout
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import socket
import subprocess
import sys
import uuid


STAGES = {"stage1": "baseline", "stage2": "stage2"}
PREVIOUS = {"stage2": "baseline"}
REGRESSION = ["tests/test_agent_llm.py", "tests/test_be04.py", "tests/test_be06.py",
              "tests/test_demo.py", "tests/test_be07.py"]


def require(condition, message):
    if not condition:
        raise ValueError(message)


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write(path, value):
    path = Path(path)
    temporary = path.with_name(path.name + ".writing")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def tree_hashes(folder):
    folder = Path(folder)
    return {p.relative_to(folder).as_posix(): digest(p)
            for p in sorted(folder.rglob("*")) if p.is_file()}


def load_manifest(path):
    path = Path(path).resolve()
    entries = read(path)["documents"]
    require(bool(entries), "Empty manifest")
    labels = set()
    for entry in entries:
        label = entry["label"]
        require(isinstance(label, str) and re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]*", label),
                "Labels must be safe directory names")
        require(label.lower() not in labels and label.upper() not in
                {"CON", "PRN", "AUX", "NUL", *[f"COM{i}" for i in range(1, 10)],
                 *[f"LPT{i}" for i in range(1, 10)]}, "Duplicate/reserved label")
        labels.add(label.lower())
        for key in ("document_path", "assets_dir"):
            entry[key] = (path.parent / entry[key]).resolve()
    return entries


def code_identity(root):
    """Actual file hashes, including uncommitted edits; HEAD alone is insufficient."""
    names = subprocess.check_output(
        ["git", "ls-files", "--cached", "--others", "--exclude-standard", "-z", "--", "app", "prompts", "tests", "scripts/experiments",
         "pyproject.toml", "uv.lock"], cwd=root).decode().split("\0")
    files = {name: digest(root / name) for name in names if name and (root / name).is_file()}
    return {"head": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=root, text=True).strip(),
            "files": files, "runner_sha256": digest(__file__)}


def verify_resume(record, identity, context, output):
    require(record["code"]["files"] == identity["files"] and
            record["code"]["runner_sha256"] == identity["runner_sha256"],
            "Product/test/runner changed: start a separate run")
    require(record["context"] == context, "Input, baseline or previous stage changed")
    for row in record["results"]:
        folder = output / row["label"]
        require(tree_hashes(folder) == row["artifacts"], "Completed artifact changed or missing")


@contextmanager
def offline_network():
    original = socket.socket.connect
    original_ex = socket.socket.connect_ex

    def guard(address):
        if isinstance(address, tuple):
            require(address[0] in ("127.0.0.1", "localhost", "::1"),
                    "External network disabled for saved-draft evaluation")

    def connect(sock, address):
        guard(address)
        return original(sock, address)

    def connect_ex(sock, address):
        guard(address)
        return original_ex(sock, address)

    socket.socket.connect, socket.socket.connect_ex = connect, connect_ex
    try:
        yield
    finally:
        socket.socket.connect, socket.socket.connect_ex = original, original_ex


def compact(text):
    return re.sub(r"\s+", "", text)


def preservation(before, after, stage, audit):
    """This migration permits pagination only; never rewrite a saved block."""
    old = [b.model_dump(mode="json") for p in before.pages for b in p.blocks]
    new = [b.model_dump(mode="json") for p in after.pages for b in p.blocks]
    require([b["block_id"] for b in old] == [b["block_id"] for b in new], "Block ID/order loss")
    require(old == new and not audit, "Unrecorded body change")
    return {"blocks": len(old), "block_ids_and_order_preserved": True,
            "original_provenance_preserved": True, "changes": []}


def previews(pdf, folder):
    import pypdfium2 as pdfium
    from PIL import Image, ImageDraw

    with pdfium.PdfDocument(pdf) as reader:
        sheet = Image.new("RGB", (1680, 610 * ((len(reader) + 3) // 4)), "#dce1e5")
        draw = ImageDraw.Draw(sheet)
        for i in range(len(reader)):
            page = reader[i]
            bitmap = page.render(scale=1.1)
            image = bitmap.to_pil().convert("RGB")
            image.save(folder / f"page-{i + 1}.png")
            image.thumbnail((400, 566))
            x, y = (i % 4) * 420 + 10, (i // 4) * 610 + 25
            sheet.paste(image, (x, y))
            draw.text((x, y - 18), str(i + 1), fill="black")
            bitmap.close()
            page.close()
        sheet.save(folder / "sheet.png")


def inspect_pdf(result, document, folder):
    from pypdf import PdfReader
    from scripts.experiments.agent_quality_comparison import layout_diagnostics

    reader = PdfReader(result.file_path, strict=True)
    require(result.layout_ok and result.actual_pages == len(document.pages) == len(reader.pages),
            "Layout/page-count failure")
    text = compact("".join(page.extract_text() or "" for page in reader.pages))
    missing = []
    for page in document.pages:
        for block in page.blocks:
            content = block.content
            strings = content.get("items", []) if block.type == "list" else [content.get("text", content.get("caption", ""))]
            if any(compact(str(s)) not in text for s in strings if s):
                missing.append(block.block_id)
    require(not missing, f"Missing PDF text: {missing}")
    previews(result.file_path, folder)
    write(folder / "document.json", document.model_dump(mode="json"))
    write(folder / "render.json", result.to_dict())
    return {"pages": len(reader.pages), "layout_ok": True,
            "target_pages_match": result.actual_pages == document.target_pages, "missing_text_blocks": missing,
            "diagnostics": layout_diagnostics(result), "pdf": str(result.file_path),
            "pdf_sha256": digest(result.file_path), "image_count": sum(len(p.images) for p in reader.pages)}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("stage", choices=STAGES)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--run-dir", type=Path, required=True)
    parser.add_argument("--baseline-dir", type=Path)
    parser.add_argument("--previous-dir", type=Path)
    parser.add_argument("--expected-template", required=True)
    parser.add_argument("--resume", action="store_true")
    parser.add_argument("--docx-label", action="append", default=[])
    parser.add_argument("--smoke-four", action="store_true")
    args = parser.parse_args(argv)
    root = Path(__file__).resolve().parents[2]
    base, manifest = args.run_dir.resolve(), args.manifest.resolve()
    output = base / STAGES[args.stage]
    previous = (args.previous_dir or base / PREVIOUS.get(args.stage, "baseline")).resolve()
    baseline = (args.baseline_dir or base / "baseline").resolve()
    entries = load_manifest(manifest)
    require(set(args.docx_label) <= {e["label"] for e in entries}, "Unknown selected label")
    require(not output.exists() or args.resume, "Output exists; use another run directory or verified --resume")
    require(not args.resume or (output / "results.json").is_file(), "No completed resume record")
    if args.stage != "stage1":
        require(previous.is_dir() and baseline.is_dir(), "Previous stage/baseline missing")
        require(output != previous and output != baseline, "Output overlaps immutable input")
    for entry in entries:
        require(not entry["document_path"].is_relative_to(output) and
                not entry["assets_dir"].is_relative_to(output), "Output overlaps manifest input")
    code = code_identity(root)
    context = {"stage": args.stage, "manifest_sha256": digest(manifest),
               "inputs": {e["label"]: {"document": digest(e["document_path"]),
                          "assets": tree_hashes(e["assets_dir"])} for e in entries},
               "previous": tree_hashes(previous) if args.stage != "stage1" else {},
               "baseline": tree_hashes(baseline) if args.stage != "stage1" else {},
               "options": {k: getattr(args, k) for k in ("expected_template", "docx_label", "smoke_four")}}
    record = read(output / "results.json") if args.resume else {"code": code, "context": context, "results": [], "executions": []}
    if args.resume:
        verify_resume(record, code, context, output)
    execution = {"id": str(uuid.uuid4()), "started_at": datetime.now(timezone.utc).isoformat(), "code": code}
    record["executions"].append(execution)
    output.mkdir(parents=True, exist_ok=args.resume)
    temp = base / "_tmp"
    temp.mkdir(exist_ok=True)
    os.environ.update(PYTHON_DOTENV_DISABLED="1", PYTHONDONTWRITEBYTECODE="1", PYTHONIOENCODING="utf-8",
                      TEMP=str(temp), TMP=str(temp))
    sys.dont_write_bytecode = True
    sys.path[:0] = [str(root), str(root / "tests")]
    os.chdir(root)
    with offline_network():
        from app.config import Settings
        from app.models import Document
        from app.services import export_render as er, layout_checks

        require(layout_checks.TEMPLATE_VERSION == args.expected_template, "Wrong product template revision")
        settings = Settings(private_runs_dir=output, db_path=output / "unused.sqlite3")
        record.update(root=str(root), renderer_module=er.__file__, template_version=layout_checks.TEMPLATE_VERSION,
                      template_fingerprint=er.template_fingerprint(), api_calls=0)
        write(output / "results.json", record)
        for entry in entries:
            label = entry["label"]
            if any(row["label"] == label for row in record["results"]):
                continue
            folder = output / label
            require(not folder.exists(), "Partial document exists; explicitly archive it before resuming")
            folder.mkdir()
            doc = Document.model_validate_json((entry["document_path"] if args.stage == "stage1" else previous / label / "document.json").read_bytes())
            before, audit = doc.model_copy(deep=True), []
            assets = {}
            asset_root = entry["assets_dir"]
            from scripts.experiments.agent_quality_comparison import load_assets
            assets = load_assets(asset_root, read(asset_root / "asset_manifest.json"))
            if args.stage == "stage2":
                prepared = er.paginate_draft(er.snapshot_from_document(doc, assets), folder / "pagination", settings)
                require(prepared.outcome == "passed", "Pagination failed")
                doc = doc.model_copy(update={"pages": prepared.pages})
            preserved = preservation(before, doc, args.stage, audit)
            result = er.render(er.snapshot_from_document(doc, assets), "pdf", folder, settings)
            row = inspect_pdf(result, doc, folder)
            row.update(label=label, execution_id=execution["id"], source_sha256=context["inputs"][label]["document"],
                       preservation=preserved, audit=audit)
            if label in args.docx_label:
                import docx
                word = er.render(er.snapshot_from_document(doc, assets), "docx", folder, settings)
                opened = docx.Document(word.file_path)
                text = compact("".join(p.text for p in opened.paragraphs))
                strings = [str(b.content.get("text", b.content.get("caption", ""))) for p in doc.pages for b in p.blocks]
                require(all(compact(s) in text for s in strings if s), "DOCX text loss")
                require(len(opened.inline_shapes) == sum(b.type == "image" for p in doc.pages for b in p.blocks), "DOCX photo loss")
                row["docx_smoke"] = {"text_and_images_preserved": True, "visual_pagination": next(c.result for c in word.checks if c.check_key == "overflow"), "layout_ok": word.layout_ok}
            row["artifacts"] = tree_hashes(folder)
            record["results"].append(row)
            write(output / "results.json", record)
            print(json.dumps({"label": label, "pages": row["pages"], "peak": row["diagnostics"]["body_peak"]}), flush=True)
        if args.smoke_four:
            from test_be07 import _fixture_snapshot
            smoke = output / ("smoke4_" + execution["id"])
            smoke.mkdir()
            result = er.render(_fixture_snapshot("4pages"), "pdf", smoke, settings)
            require(result.layout_ok and result.actual_pages == 4, "Four-page smoke failed")
            previews(result.file_path, smoke)
            write(smoke / "render.json", result.to_dict())
        import pytest
        log = output / ("regression_" + execution["id"] + ".log")
        with log.open("w", encoding="utf-8") as stream, redirect_stdout(stream), redirect_stderr(stream):
            status = pytest.main(["-q", "-p", "no:cacheprovider", "--import-mode=importlib",
                                  "--basetemp=" + str(temp / ("pytest_" + execution["id"])), *REGRESSION])
        write(output / "regression_exit.json", {"exit_code": int(status), "log": str(log),
              "execution_id": execution["id"], "code": code, "reused": False})
        print(log.read_text(encoding="utf-8")[-2400:], flush=True)
        require(status == 0, "Regression failed")
        require(code_identity(root) == code, "Code changed during execution")
        for entry in entries:
            expected = context["inputs"][entry["label"]]
            require(digest(entry["document_path"]) == expected["document"] and
                    tree_hashes(entry["assets_dir"]) == expected["assets"], "Saved source changed")
        if args.stage != "stage1":
            require(tree_hashes(baseline) == context["baseline"] and tree_hashes(previous) == context["previous"], "Baseline/previous stage changed")
        write(output / "ready_for_review.json", {"execution_id": execution["id"], "regression_exit": 0,
              "documents": len(record["results"]), "visual_review": "pending", "temporary_dir": str(temp)})
        print("Ready for visual review; no Git write or automatic cleanup performed.", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
