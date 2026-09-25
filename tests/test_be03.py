"""BE-03 확인: 형식 7종 읽기, 스캔 PDF partial, 암호·손상·빈 파일, locator 규칙, 이미지 asset, Job 진행, v1→v2 마이그레이션.

테스트 파일은 전부 여기서 만든다(가짜 데이터). 실제 회사 자료는 레포에 넣지 않는다.
실행: uv run pytest
"""
from __future__ import annotations

import io
import json
import sqlite3

import pytest
from fastapi.testclient import TestClient

from app import create_app
from app.config import Settings
from app.db import connect, init_db
from app.parsers import parse
from app.parsers.text import split_lines

from pathlib import Path

FIXTURES = Path(__file__).parent / "fixtures"
BRIEF = {"purpose": "테스트", "emphasis": [], "direction": "balanced", "target_pages": 4, "photo_preference": "balanced"}


# ---------------- 가짜 파일 생성 ----------------

def make_pdf(pages: list[str | None]) -> bytes:
    """페이지별 텍스트(None이면 글자 없는 빈 페이지 = 스캔본 흉내)로 최소 PDF를 만든다. ASCII만."""
    objs: list[bytes] = []
    n_pages = len(pages)
    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(n_pages))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    font_no = 3 + n_pages * 2
    for i, text in enumerate(pages):
        page_no, content_no = 3 + i * 2, 4 + i * 2
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] /Contents {content_no} 0 R "
                    f"/Resources << /Font << /F1 {font_no} 0 R >> >> >>".encode())
        stream = f"BT /F1 12 Tf 20 150 Td ({text}) Tj ET".encode() if text else b""
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for no, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{no} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_encrypted_pdf() -> bytes:
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(["secret"]))))
    writer.encrypt("pw", algorithm="RC4-128")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pptx(slide_texts: list[list[str]], table_on_slide: int | None = None) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for s, texts in enumerate(slide_texts, start=1):
        slide = prs.slides.add_slide(blank)
        for i, text in enumerate(texts):
            box = slide.shapes.add_textbox(Inches(1), Inches(0.5 + i), Inches(6), Inches(0.8))
            box.text_frame.text = text
        if table_on_slide == s:
            shape = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(5), Inches(1))
            shape.table.cell(0, 0).text = "항목"
            shape.table.cell(0, 1).text = "값"
            shape.table.cell(1, 0).text = "공정 수"
            shape.table.cell(1, 1).text = "2개"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_png(width: int = 12, height: int = 8) -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (width, height), (200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


# ---------------- 파서 단위 ----------------

def test_txt_fixture_line_locators():
    data = (FIXTURES / "mock_source_a.txt").read_bytes()
    r = parse(data, ".txt")
    assert r.status == "complete" and r.text_available
    lines = split_lines(data.decode("utf-8-sig"))
    assert [s.locator for s in r.segments] == [{"line_start": i, "line_end": i} for i in range(1, 6)]
    assert r.segments[0].text == "회사명: 테스트 회사" and r.segments[4].text == lines[4].strip()


def test_txt_empty_and_bad_encoding():
    assert parse(b"   \n\n", ".txt").status == "failed"
    assert parse(b"   \n\n", ".txt").warnings[0]["code"] == "NO_USABLE_TEXT"
    r = parse("한글".encode("cp949"), ".txt")
    assert r.status == "failed" and r.warnings[0]["code"] == "UNSUPPORTED_ENCODING"


def test_md_keeps_markup_and_skips_blank_lines_but_keeps_numbers():
    r = parse("# 제목\n\n- 항목 **강조**\n".encode(), ".md")
    assert [(s.locator["line_start"], s.text) for s in r.segments] == [(1, "# 제목"), (3, "- 항목 **강조**")]


def test_pdf_text_pages():
    r = parse(make_pdf(["Hello page one", "Second page"]), ".pdf")
    assert r.status == "complete" and r.text_available
    assert [s.locator for s in r.segments] == [{"page": 1}, {"page": 2}]
    assert "Hello page one" in r.segments[0].text


def test_pdf_scanned_is_partial_with_image_only_warning():
    r = parse(make_pdf([None, None]), ".pdf")
    assert r.status == "partial" and r.text_available is False and r.segments == []
    assert r.warnings[0]["code"] == "IMAGE_ONLY" and "글자를 읽지 못함" in r.warnings[0]["message"]
    assert r.warnings[0]["locator"] == {"pages": [1, 2]}


def test_pdf_mixed_pages_partial():
    r = parse(make_pdf(["text", None, "more"]), ".pdf")
    assert r.status == "partial" and [s.locator["page"] for s in r.segments] == [1, 3]
    assert r.warnings[0]["locator"] == {"pages": [2]}


def test_pdf_encrypted_and_corrupt():
    r = parse(make_encrypted_pdf(), ".pdf")
    assert r.status == "failed" and r.warnings[0]["code"] == "ENCRYPTED"
    r = parse(b"%PDF-1.4 garbage", ".pdf")
    assert r.status == "failed" and r.warnings[0]["code"] == "FILE_CORRUPT"


def test_docx_paragraphs_and_table_cells():
    r = parse(make_docx(["회사명: 예시", "", "사업: 표면처리"], [["항목", "값"], ["공정 수", "2개"]]), ".docx")
    assert r.status == "complete"
    locs = [s.locator for s in r.segments]
    assert {"paragraph": 1} in locs and {"paragraph": 3} in locs and {"paragraph": 2} not in locs
    assert {"table": 1, "row": 2, "col": 2} in locs
    assert next(s.text for s in r.segments if s.locator == {"table": 1, "row": 2, "col": 2}) == "2개"


def test_docx_corrupt_and_empty():
    r = parse(b"not a zip at all", ".docx")
    assert r.status == "failed" and r.warnings[0]["code"] == "FILE_CORRUPT"
    r = parse(make_docx([""]), ".docx")
    assert r.status == "failed" and r.warnings[0]["code"] == "NO_USABLE_TEXT"


def test_pptx_20_slides_textboxes_and_table():
    slides = [[f"슬라이드 {i} 제목", f"본문 {i}"] for i in range(1, 21)]
    r = parse(make_pptx(slides, table_on_slide=12), ".pptx")
    assert r.status == "complete"
    assert {s.locator["slide"] for s in r.segments} == set(range(1, 21))
    assert {"slide": 12, "shape": 1} in [s.locator for s in r.segments]
    table_cells = [s for s in r.segments if s.locator.get("slide") == 12 and "row" in s.locator]
    assert len(table_cells) == 4 and table_cells[-1].locator == {"slide": 12, "shape": 3, "row": 2, "col": 2}
    assert table_cells[-1].text == "2개"


def test_png_is_image_only_with_size():
    r = parse(make_png(12, 8), ".png")
    assert r.status == "complete" and r.image_available and not r.text_available
    assert (r.width, r.height) == (12, 8) and r.warnings[0]["code"] == "IMAGE_ONLY"
    r = parse(make_png(), ".jpg")  # PNG 내용에 .jpg 확장자
    assert r.status == "failed" and r.warnings[0]["code"] == "FILE_CORRUPT"
    assert parse(b"\x89PNG broken", ".png").status == "failed"


# ---------------- API 통합 ----------------

@pytest.fixture
def settings(tmp_path):
    return Settings(private_runs_dir=tmp_path / "runs", db_path=tmp_path / "runs" / "t.sqlite3")


@pytest.fixture
def client(settings):
    return TestClient(create_app(settings))


def _session(client: TestClient) -> str:
    r = client.post("/api/v1/sessions", json={"brief": BRIEF})
    assert r.status_code == 201
    return r.json()["session_id"]


def _upload(client: TestClient, sid: str, *files: tuple[str, bytes], kind: str | None = None) -> dict:
    data = {"kind": kind} if kind else None
    r = client.post(f"/api/v1/sessions/{sid}/sources", data=data,
                    files=[("files", (name, io.BytesIO(content))) for name, content in files])
    assert r.status_code == 202, r.text
    return r.json()


def test_upload_then_background_read_updates_status_and_job(client, settings):
    sid = _session(client)
    up = _upload(client, sid, ("a.txt", (FIXTURES / "mock_source_a.txt").read_bytes()),
                 ("scan.pdf", make_pdf([None])), ("deck.pptx", make_pptx([["한 장"]])))
    assert all(i["parse_status"] == "queued" for i in up["items"])  # 202 시점에는 아직 대기
    # TestClient는 응답 뒤 백그라운드 작업까지 돌린 다음 돌아온다.
    job = client.get(f"/api/v1/sessions/{sid}/jobs/{up['job_id']}").json()
    assert job["status"] == "succeeded" and job["progress"] == {"stage": "done", "message": None}
    assert job["result_ref"] == {"type": "sources", "source_ids": [i["source_id"] for i in up["items"]]}
    items = client.get(f"/api/v1/sessions/{sid}/sources").json()["items"]
    by_name = {i["name"]: i for i in items}
    txt = by_name["a.txt"]
    assert txt["parse_status"] == "complete" and txt["text_available"] and len(txt["usable_segment_ids"]) == 5
    scan = by_name["scan.pdf"]
    assert scan["parse_status"] == "partial" and not scan["text_available"] and scan["usable_segment_ids"] == []
    assert scan["warnings"][0]["code"] == "IMAGE_ONLY"
    assert by_name["deck.pptx"]["parse_status"] == "complete"
    with connect(settings.db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM segments WHERE session_id=?", (sid,)).fetchone()[0] == 6
        path = conn.execute("SELECT stored_path FROM sources WHERE source_id=?", (txt["source_id"],)).fetchone()[0]
        assert path == f"{sid}/{txt['source_id']}.txt"  # private_runs 기준 상대경로


def test_char_limit_marks_partial_not_silent_truncation(tmp_path):
    client = TestClient(create_app(Settings(private_runs_dir=tmp_path / "runs", db_path=tmp_path / "runs" / "t.sqlite3",
                                            max_source_chars=60)))
    sid = _session(client)
    long_txt = "\n".join(f"{i:02d} " + "가" * 20 for i in range(1, 8)).encode()  # 7행 × 23자 > 60
    up = _upload(client, sid, ("long.txt", long_txt))
    item = client.get(f"/api/v1/sessions/{sid}/sources").json()["items"][0]
    assert item["parse_status"] == "partial" and len(item["usable_segment_ids"]) == 2
    assert item["warnings"][0]["code"] == "TEXT_LIMIT" and item["warnings"][0]["locator"] == {"line_start": 2, "line_end": 2}


def test_failed_file_does_not_block_others(client):
    sid = _session(client)
    _upload(client, sid, ("bad.docx", b"garbage"), ("ok.txt", b"hello"))
    items = client.get(f"/api/v1/sessions/{sid}/sources").json()["items"]
    assert [i["parse_status"] for i in items] == ["failed", "complete"]
    assert items[0]["warnings"][0]["code"] == "FILE_CORRUPT"


def test_image_upload_creates_asset_and_serves_bytes(client):
    sid = _session(client)
    png = make_png(5, 7)
    up = _upload(client, sid, ("photo.png", png))
    assert up["items"][0]["kind"] == "photo"  # kind 미지정 이미지는 photo
    item = client.get(f"/api/v1/sessions/{sid}/sources").json()["items"][0]
    assert item["image_available"] and not item["text_available"] and len(item["asset_ids"]) == 1
    r = client.get(f"/api/v1/sessions/{sid}/assets/{item['asset_ids'][0]}")
    assert r.status_code == 200 and r.headers["content-type"].startswith("image/png") and r.content == png
    assert r.headers["cache-control"] == "private, no-store"


def test_asset_not_visible_to_other_owner_or_after_delete(client, settings):
    sid = _session(client)
    _upload(client, sid, ("photo.png", make_png()))
    item = client.get(f"/api/v1/sessions/{sid}/sources").json()["items"][0]
    aid = item["asset_ids"][0]
    other = TestClient(create_app(settings))
    other.post("/api/v1/sessions", json={"brief": BRIEF})
    assert other.get(f"/api/v1/sessions/{sid}/assets/{aid}").status_code == 404
    client.delete(f"/api/v1/sessions/{sid}/sources/{item['source_id']}", params={"expected_input_revision": 1})
    assert client.get(f"/api/v1/sessions/{sid}/assets/{aid}").status_code == 404
    assert client.get(f"/api/v1/sessions/{sid}/sources").json()["items"] == []


def test_delete_source_removes_segments(client, settings):
    sid = _session(client)
    up = _upload(client, sid, ("a.txt", b"one\ntwo"))
    src = up["items"][0]["source_id"]
    with connect(settings.db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM segments WHERE source_id=?", (src,)).fetchone()[0] == 2
    client.delete(f"/api/v1/sessions/{sid}/sources/{src}", params={"expected_input_revision": 1})
    with connect(settings.db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM segments WHERE source_id=?", (src,)).fetchone()[0] == 0


def test_stale_jobs_failed_on_startup(settings):
    client = TestClient(create_app(settings))
    sid = _session(client)
    with connect(settings.db_path) as conn:
        conn.execute("INSERT INTO jobs (job_id, session_id, kind, status, progress_json, created_at, updated_at) "
                     "VALUES ('job_stale', ?, 'read', 'running', '{\"stage\": \"reading\", \"message\": \"1/2\"}', 'x', 'x')",
                     (sid,))
    client2 = TestClient(create_app(settings))
    client2.cookies = client.cookies
    job = client2.get(f"/api/v1/sessions/{sid}/jobs/job_stale").json()
    assert job["status"] == "failed" and job["error"]["code"] == "SERVICE_TEMPORARY_FAILURE"


# ---------------- 마이그레이션 ----------------

V1_SOURCES = """
CREATE TABLE sessions (session_id TEXT PRIMARY KEY, owner_id TEXT NOT NULL, status TEXT NOT NULL,
  input_revision INTEGER NOT NULL, brief_json TEXT NOT NULL, selected_source_ids TEXT NOT NULL,
  created_at TEXT NOT NULL, last_activity_at TEXT NOT NULL, expires_at TEXT NOT NULL, closed_at TEXT, cleanup_status TEXT);
CREATE TABLE sources (source_id TEXT PRIMARY KEY, session_id TEXT NOT NULL REFERENCES sessions(session_id),
  source_version INTEGER NOT NULL, scope TEXT NOT NULL, name TEXT NOT NULL, mime_type TEXT NOT NULL,
  size_bytes INTEGER NOT NULL, kind TEXT NOT NULL, parse_status TEXT NOT NULL, text_available INTEGER NOT NULL,
  image_available INTEGER NOT NULL, stored_path TEXT NOT NULL, content_hash TEXT NOT NULL, created_at TEXT NOT NULL,
  expires_at TEXT, deleted_at TEXT);
"""


def test_v1_db_migrates_to_v2_keeping_rows(tmp_path):
    runs = tmp_path / "runs"
    runs.mkdir()
    db = runs / "app.sqlite3"
    with sqlite3.connect(db) as conn:
        conn.executescript(V1_SOURCES)
        conn.execute("INSERT INTO sessions VALUES ('sess_1','own_1','active',1,'{}','[]','t','t','t',NULL,NULL)")
        conn.execute("INSERT INTO sources VALUES ('src_1','sess_1',1,'session','a.txt','text/plain',3,'other',"
                     "'queued',0,0,?, 'h','t',NULL,NULL)", (str(runs / "sess_1" / "src_1.txt"),))
    init_db(db, runs)
    with sqlite3.connect(db) as conn:
        conn.row_factory = sqlite3.Row
        from app.db import SCHEMA_VERSION
        assert conn.execute("PRAGMA user_version").fetchone()[0] == SCHEMA_VERSION
        row = conn.execute("SELECT * FROM sources").fetchone()
        assert row["stored_path"] == "sess_1/src_1.txt" and row["warnings_json"] == "[]"
        cols = {r[1]: r[3] for r in conn.execute("PRAGMA table_info(sources)")}  # name -> notnull
        assert cols["session_id"] == 0
        # 등록 자료(세션 없음) 삽입 가능, 세션 자료에 세션 없음은 CHECK로 거부
        conn.execute("INSERT INTO sources (source_id, session_id, source_version, scope, name, mime_type, size_bytes, kind, "
                     "parse_status, text_available, image_available, stored_path, content_hash, created_at) "
                     "VALUES ('src_reg', NULL, 1, 'registered', 'r.txt', 'text/plain', 1, 'company', 'queued', 0, 0, "
                     "'registered/src_reg.txt', 'h', 't')")
        with pytest.raises(sqlite3.IntegrityError):
            conn.execute("INSERT INTO sources (source_id, session_id, source_version, scope, name, mime_type, size_bytes, "
                         "kind, parse_status, text_available, image_available, stored_path, content_hash, created_at) "
                         "VALUES ('src_bad', NULL, 1, 'session', 'x', 'y', 1, 'other', 'queued', 0, 0, 'p', 'h', 't')")
        for table in ("segments", "assets", "preflights", "documents", "document_revisions"):
            assert conn.execute("SELECT 1 FROM sqlite_master WHERE name=?", (table,)).fetchone()
        assert "input_revision" in [r[1] for r in conn.execute("PRAGMA table_info(jobs)")]  # v3 컬럼
    # 두 번째 init_db는 아무것도 바꾸지 않는다
    init_db(db, runs)
