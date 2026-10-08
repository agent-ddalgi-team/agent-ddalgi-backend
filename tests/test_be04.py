"""BE-04 확인: 사전 점검·초안 Job, 오래된 입력·미확인·중복·만료 처리, AI 결과 검사, 문서 버전 구조, llm 미연결.

mock Agent 결과는 가짜 회사("예시 회사") 기준. 실제 회사 정보 없음. 실제 LLM 호출 없음.
"""
from __future__ import annotations

import io
import json
from dataclasses import replace

import pytest
from fastapi.testclient import TestClient

from app import create_app
from app.agent_bridge import AgentError, AnalyzeResult, DraftResult
from app.agent_mock import MockAgent
from app.config import Settings
from app.db import connect, init_orm_db
from app.models import Block, EvidenceRef, Page
from app.services import ai_jobs

BRIEF = {"purpose": "테스트", "emphasis": [], "direction": "balanced", "target_pages": 4, "photo_preference": "balanced"}
SOURCE_A = "회사명: 예시 회사\n회사 개요: 예시용 기업입니다.\n사업 분야: 예시 사업 A\n공정 수: 2개\n납기 표현: 빠른 납기\n".encode()
SOURCE_B = "회사명: 다른 예시 회사\n".encode()

# Shared synthetic source; scoring requirements stay outside Agent input.
# These fixtures do not evaluate MockAgent as semantic AI quality.
CUSTOMER_QUALITY_SOURCE = """회사명: 예시정공
회사 개요: 예시정공은 시험용 금속 부품을 가공하는 가상 제조업체입니다.
사업 분야: 시험용 금속 부품 가공과 시편 검사 지원.
제품 및 서비스: 시험용 브래킷과 커버를 가공합니다.
공정: 절삭과 연마를 수행합니다. 공정 순서는 지정하지 않습니다.
설비: 3축 CNC 가공기와 표면 거칠기 측정기를 사용합니다.
가공 능력: 알루미늄 시편에 한해 최대 가공 길이 200mm입니다. 다른 소재의 가공 한계는 확인 필요입니다.
인증: ISO 9001 인증의 적용 범위는 시험 부품 제조이며 유효기간은 2025년부터 2027년까지입니다.
시험: 알루미늄 시험시편의 중성 염수분무 시험 시간은 168시간입니다. 시험시편 결과이며 양산 제품의 성능 보증이 아닙니다.
적용 사례: 시험용 브래킷은 연구용 고정 지그에 적용했습니다. 고객명과 판매 수량은 공개하지 않습니다.
연혁: 2024년 시범 생산을 시작했습니다.
"""
CUSTOMER_QUALITY_CASES = {
    "quality": {
        "purpose": "인증 적용 범위와 시험 근거 중심 회사 소개",
        "audience": "품질 담당자", "emphasis": ["품질관리", "인증·특허"],
        "required_fields": ["certifications", "technology"],
        "probes": {"certification": ["ISO", "9001", "시험", "제조", "2025", "2027"],
                   "test": ["알루미늄", "시편", "염수", "168", "양산", "보증"]},
    },
    "production": {
        "purpose": "공정·설비·소재와 기술 검토사항 중심 회사 소개",
        "audience": "생산기술 담당자", "emphasis": ["공정 역량", "기술"],
        "required_fields": ["processes", "capabilities"],
        "probes": {"process": ["절삭", "연마"], "equipment": ["CNC", "거칠기"],
                   "capability": ["알루미늄", "시편", "200"]},
    },
    "customer": {
        "purpose": "주요 서비스와 관련 적용 사례 중심 신규 고객 소개",
        "audience": "신규 고객", "emphasis": ["제품·서비스", "적용 사례"],
        "required_fields": ["products_services", "customers_markets"],
        "probes": {"service": ["브래킷", "커버"], "application": ["브래킷", "연구", "지그"]},
    },
}

# Live evaluation controls: explicit performance/completion must survive too.
CUSTOMER_TEST_STATE_CONTROLS = {
    "conditions": "알루미늄 시험시편의 중성 염수분무 시험 시간은168시간입니다. 시험시편 결과이며 양산 제품의 성능 보증이 아닙니다.",
    "performed": "알루미늄 시험시편에 중성 염수분무 시험을168시간 실시했습니다. 시험시편 결과이며 양산 제품의 성능 보증이 아닙니다.",
    "completed": "알루미늄 시험시편의 중성 염수분무 시험을168시간 실시하고 완료했습니다. 시험시편 결과이며 양산 제품의 성능 보증이 아닙니다.",
}

# Held-out company: identical core claims in short text and a longer DOCX table.
LONG_TABLE_QUALITY_ROWS = [
    ("회사명", "예시유체"),
    ("회사 개요", "예시유체는 연구용 냉각 모듈을 조립하고 시험하는 가상 기업입니다."),
    ("제품 및 서비스", "연구용 냉각 모듈 CM-20의 조립과 시험시편 검사를 제공합니다."),
    ("공정", "부품 세척과 모듈 조립을 수행합니다. 공정 순서는 지정하지 않습니다."),
    ("설비", "유량 시험대와 압력 센서를 사용합니다."),
    ("가공 능력", "알루미늄 시험시편의 최대 가공 길이는 200mm입니다. 스테인리스 시험시편은 100mm입니다."),
    ("시험 조건", "CM-20 시험시편의 유량은 35L/min, 최대 압력은 0.6MPa입니다. 물 온도 20±2°C와 회전수 1,450rpm 조건의 시험시편 수치이며 양산 제품의 성능 보증이 아닙니다."),
    ("인증", "ISO 9001의 적용 범위는 연구용 냉각 모듈 조립입니다. 유효기간은 2025년 1월 1일부터 2027년 12월 31일까지입니다."),
    ("납기", "표준 주문은 도면 승인 후 영업일 7일입니다. 시제품 주문은 사양 확정 후 영업일 20일입니다."),
    ("연혁", "2009년 법인을 설립했고 2018년 연구용 냉각 모듈의 시범 생산을 시작했습니다."),
]
LONG_TABLE_QUALITY_BACKGROUND = [
    "자료 배경: 예시유체의 연구용 모듈 작업 기록에는 도면 검토와 부품 상태 확인 내용을 함께 남깁니다.",
    "자료 배경: 시험시편의 소재와 시험 조건은 작업 기록에 구분해 적으며 다른 제품의 성능으로 일반화하지 않습니다.",
    "자료 배경: 고객과 협의할 때 표준 주문과 시제품 주문의 승인 조건을 구분하고 변경된 사양은 다시 확인합니다.",
    "자료 배경: 이 자료는 가상 기업의 소개서 시험 자료이며 고객명이나 판매 실적을 제시하지 않습니다.",
]
LONG_TABLE_QUALITY_PROBES = {
    "materials": ["알루미늄", "200", "스테인리스", "100"],
    "test_scope": ["35", "0.6", "20", "2", "1,450", "시험시편", "양산", "보증"],
    "certification": ["ISO", "9001", "조립", "2025", "2027"],
    "delivery": ["도면", "승인", "영업일", "7", "사양", "확정", "20"],
}

MULTISOURCE_TIME_CASES = {
    "different_dates": [
        "자료 A\n회사명: 예시표면\n회사 개요: 예시표면은 시험용 금속 부품의 표면처리를 수행하는 가상 기업입니다.\n설비 현황: 2022년 10월 1일 기준 표면처리 라인은 3개입니다.\n",
        "자료 B\n회사명: 예시표면\n회사 개요: 예시표면은 시험용 금속 부품의 표면처리를 수행하는 가상 기업입니다.\n설비 현황: 2026년 10월 1일 기준 표면처리 라인은 4개입니다.\n",
    ],
    "same_date_conflict": [
        "자료 A\n회사명: 예시표면\n회사 개요: 예시표면은 시험용 금속 부품의 표면처리를 수행하는 가상 기업입니다.\n설비 현황: 2026년 10월 1일 기준 표면처리 라인은 3개입니다.\n",
        "자료 B\n회사명: 예시표면\n회사 개요: 예시표면은 시험용 금속 부품의 표면처리를 수행하는 가상 기업입니다.\n설비 현황: 2026년 10월 1일 기준 표면처리 라인은 4개입니다.\n",
    ],
}


def long_table_quality_source(fmt):
    if fmt == "txt":
        return "\n".join(f"{key}: {value}" for key, value in LONG_TABLE_QUALITY_ROWS).encode()
    from docx import Document
    document = Document()
    document.add_heading("예시유체 연구용 모듈 소개 자료", 0)
    for index in range(36):
        document.add_paragraph(LONG_TABLE_QUALITY_BACKGROUND[index % 4])
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "항목", "내용 및 적용 조건"
    for key, value in LONG_TABLE_QUALITY_ROWS:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, value
    for index in range(36):
        document.add_paragraph(LONG_TABLE_QUALITY_BACKGROUND[index % 4])
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def customer_quality_brief(case_id):
    case = CUSTOMER_QUALITY_CASES[case_id]
    return {"purpose": case["purpose"], "audience": case["audience"],
            "emphasis": list(case["emphasis"]), "required_fields": list(case["required_fields"]),
            "direction": "balanced", "target_pages": 1, "photo_preference": "none",
            "target_company": "예시정공"}


@pytest.fixture
def settings(tmp_path):
    return Settings(private_runs_dir=tmp_path / "runs", db_path=tmp_path / "runs" / "t.sqlite3")


@pytest.fixture
def client(settings):
    return TestClient(create_app(settings))


def _session(client) -> str:
    return client.post("/api/v1/sessions", json={"brief": BRIEF}).json()["session_id"]


def _upload(client, sid, *files, kind=None) -> list[str]:
    r = client.post(f"/api/v1/sessions/{sid}/sources", data={"kind": kind} if kind else None,
                    files=[("files", (n, io.BytesIO(c))) for n, c in files])
    assert r.status_code == 202, r.text
    return [i["source_id"] for i in r.json()["items"]]


def _select(client, sid, source_ids, expected=1) -> int:
    r = client.patch(f"/api/v1/sessions/{sid}/inputs",
                     json={"expected_input_revision": expected, "selected_source_ids": source_ids})
    assert r.status_code == 200, r.text
    return r.json()["input_revision"]


def _preflight(client, sid, rev) -> dict:
    r = client.post(f"/api/v1/sessions/{sid}/preflights", json={"expected_input_revision": rev})
    assert r.status_code == 202, r.text
    job = client.get(f"/api/v1/sessions/{sid}/jobs/{r.json()['job_id']}").json()
    assert job["status"] == "succeeded", job
    return client.get(f"/api/v1/sessions/{sid}/preflights/{job['result_ref']['preflight_id']}").json()


def _png() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (4, 4), (0, 0, 0)).save(buf, format="PNG")
    return buf.getvalue()


# ---------------- 사전 점검 ----------------

def test_preflight_facts_issues_and_can_generate(client):
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A), ("photo.png", _png()))
    rev = _select(client, sid, src)
    pf = _preflight(client, sid, rev)
    assert pf["input_revision"] == rev and pf["can_generate"] is True and pf["confirmed_at"] is None
    assert pf["usable_source_ids"] == [src[0]]  # 이미지는 텍스트 근거가 아님
    by_key = {f["field_key"]: f for f in pf["facts"]}
    assert len(by_key) == 14
    assert by_key["company_name"]["status"] == "supported" and by_key["company_name"]["value"] == "예시 회사"
    ev = by_key["company_name"]["evidence_refs"][0]
    assert ev["source_id"] == src[0] and ev["segment_id"].startswith("seg_") and ev["locator"] == {"line_start": 1, "line_end": 1}
    assert by_key["lead_time"]["status"] == "needs_confirmation"
    assert by_key["certifications"]["status"] == "missing" and by_key["certifications"]["evidence_refs"] == []
    codes = {i["code"]: i for i in pf["issues"]}
    assert codes["LEAD_TIME_UNQUANTIFIED"]["severity"] == "warning" and codes["LEAD_TIME_UNQUANTIFIED"]["fact_ids"] == [by_key["lead_time"]["fact_id"]]
    assert codes["IMAGE_ONLY_SOURCE"]["source_ids"] == [src[1]]
    assert "REQUIRED_MISSING" not in codes
    assert pf["recommendations"]["suggested_pages"] == 4


def test_preflight_conflict_and_required_missing(client):
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A), ("b.txt", SOURCE_B))
    pf = _preflight(client, sid, _select(client, sid, src))
    name = next(f for f in pf["facts"] if f["field_key"] == "company_name")
    assert name["status"] == "conflict" and name["value"] is None and len(name["alternatives"]) == 2
    assert any(i["code"] == "VALUE_CONFLICT" and i["severity"] == "blocker" for i in pf["issues"])
    sid2 = _session(client)
    src2 = _upload(client, sid2, ("only.txt", "사업 분야: 뭔가".encode()))
    pf2 = _preflight(client, sid2, _select(client, sid2, src2))
    assert sum(1 for i in pf2["issues"] if i["code"] == "REQUIRED_MISSING") == 2


def test_preflight_without_text_sources_cannot_generate(client):
    sid = _session(client)
    src = _upload(client, sid, ("photo.png", _png()))
    pf = _preflight(client, sid, _select(client, sid, src))
    assert pf["can_generate"] is False and pf["usable_source_ids"] == []
    r = client.post(f"/api/v1/sessions/{sid}/drafts",
                    json={"preflight_id": pf["preflight_id"], "input_revision": 2, "confirmed": True})
    assert r.status_code == 422 and r.json()["error"]["code"] == "NO_USABLE_TEXT"


def test_preflight_stale_revision_409_and_duplicate_returns_same_job(client, settings):
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A))
    rev = _select(client, sid, src)
    r = client.post(f"/api/v1/sessions/{sid}/preflights", json={"expected_input_revision": rev - 1})
    assert r.status_code == 409 and r.json()["error"]["code"] == "INPUT_REVISION_CONFLICT"
    # 진행 중인 같은 버전의 작업이 있으면 그 job을 돌려준다
    with connect(settings.db_path) as conn:
        conn.execute("INSERT INTO jobs (job_id, session_id, kind, status, progress_json, input_revision, created_at, updated_at) "
                     "VALUES ('job_running', ?, 'preflight', 'running', '{\"stage\":\"analyzing\",\"message\":null}', ?, 'x', 'x')",
                     (sid, rev))
    r = client.post(f"/api/v1/sessions/{sid}/preflights", json={"expected_input_revision": rev})
    assert r.status_code == 202 and r.json() == {"job_id": "job_running", "status": "running", "kind": "preflight",
                                                 "session_id": sid, "created_at": "x"}


def test_preflight_job_fails_if_input_changed_while_running(client, settings, monkeypatch):
    """Job이 도는 사이 입력이 바뀌면 결과를 버린다(오래된 점검 사용 금지, BR-05)."""
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A))
    rev = _select(client, sid, src)
    original = MockAgent.analyze

    async def analyze_then_bump(self, request):
        with connect(settings.db_path) as conn:
            conn.execute("UPDATE sessions SET input_revision=input_revision+1 WHERE session_id=?", (sid,))
        return await original(self, request)

    monkeypatch.setattr(MockAgent, "analyze", analyze_then_bump)
    r = client.post(f"/api/v1/sessions/{sid}/preflights", json={"expected_input_revision": rev})
    job = client.get(f"/api/v1/sessions/{sid}/jobs/{r.json()['job_id']}").json()
    assert job["status"] == "failed" and job["error"]["code"] == "INPUT_REVISION_CONFLICT"
    with connect(settings.db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM preflights WHERE session_id=?", (sid,)).fetchone()[0] == 0


@pytest.fixture(params=[9, 10], ids=["legacy-v9", "orm-v10"])
def retry_api(request, tmp_path):
    settings = Settings(private_runs_dir=tmp_path / "runs", db_path=tmp_path / "runs" / "retry.sqlite3",
                        agent_mode="mock", cleanup_sweep_interval_s=0)
    if request.param == 10:
        init_orm_db(settings.db_path, settings.private_runs_dir)
    with TestClient(create_app(settings)) as client:
        yield client, settings


@pytest.mark.parametrize("kind", ["preflight", "draft"])
@pytest.mark.parametrize("terminal", ["succeeded", "failed"])
def test_active_job_join_key_replays_after_finish_and_explicit_retry(retry_api, monkeypatch, kind, terminal):
    """진행 중 작업을 받은 요청도 첫 응답을 기억한다. 실패 재시도는 새 키로 명시한다."""
    client, settings = retry_api
    sid = _session(client)
    source_ids = _upload(client, sid, ("fake.txt", SOURCE_A))
    rev = _select(client, sid, source_ids)
    base = f"/api/v1/sessions/{sid}"
    path = f"{base}/{kind}s"
    body = {"expected_input_revision": rev}
    if kind == "draft":
        pf = _preflight(client, sid, rev)
        body = {"preflight_id": pf["preflight_id"], "input_revision": rev, "confirmed": True}
    accepted_body = dict(body)

    runner_name = f"run_{kind}_job"
    runner = getattr(ai_jobs, runner_name)
    scheduled = []
    monkeypatch.setattr(ai_jobs, runner_name, lambda *args: scheduled.append(args))
    first = client.post(path, json=body, headers={"Idempotency-Key": "original"})
    joined = client.post(path, json=body, headers={"Idempotency-Key": "joined"})
    assert first.status_code == joined.status_code == 202
    assert joined.json() == first.json()
    assert len(scheduled) == 1

    method_name = "analyze" if kind == "preflight" else "draft"
    original_method = getattr(MockAgent, method_name)
    calls = []

    async def finish(self, request):
        calls.append(request)
        if terminal == "failed":
            raise AgentError("AI_RATE_LIMIT", "가짜 일시 실패", retryable=True)
        return await original_method(self, request)

    monkeypatch.setattr(MockAgent, method_name, finish)
    runner(*scheduled[0])
    monkeypatch.setattr(ai_jobs, runner_name, runner)
    job_url = f"{base}/jobs/{first.json()['job_id']}"
    job = client.get(job_url).json()
    assert job["status"] == terminal

    replay = client.post(path, json=body, headers={"Idempotency-Key": "joined"})
    assert replay.status_code == 202 and replay.json() == joined.json()
    changed_body = dict(body)
    changed_body["expected_input_revision" if kind == "preflight" else "input_revision"] = rev + 1
    conflict = client.post(path, json=changed_body, headers={"Idempotency-Key": "joined"})
    assert conflict.status_code == 409 and conflict.json()["error"]["code"] == "IDEMPOTENCY_KEY_CONFLICT"
    assert len(calls) == 1
    with connect(settings.db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM jobs WHERE session_id=? AND kind=?", (sid, kind)).fetchone()[0] == 1
        assert conn.execute("PRAGMA foreign_key_check").fetchall() == []

    # 실패는 이전 작업에 남는다. 자료와 선택은 보존하고 새 요청으로 다시 진행할 수 있다.
    if terminal == "failed":
        monkeypatch.setattr(MockAgent, method_name, original_method)
        session = client.get(base).json()
        assert session["selected_source_ids"] == source_ids and session["input_revision"] == rev
        assert session["document_summary"] is None
        if kind == "draft":
            latest = _preflight(client, sid, rev)
            body = {**body, "preflight_id": latest["preflight_id"]}
        retry = client.post(path, json=body, headers={"Idempotency-Key": "explicit-retry"})
        assert retry.status_code == 202 and retry.json()["job_id"] != first.json()["job_id"]
        assert client.get(f"{base}/jobs/{retry.json()['job_id']}").json()["status"] == "succeeded"
        assert client.get(job_url).json() == job
        assert client.post(path, json=body, headers={"Idempotency-Key": "explicit-retry"}).json() == retry.json()

    # 새 합류 키도 세션에 묶여 종료 뒤에는 캐시 응답을 노출하지 않는다.
    assert client.delete(base).status_code == 200
    gone = client.post(path, json=accepted_body,
                       headers={"Idempotency-Key": "joined"})
    assert gone.status_code == 410 and gone.json()["error"]["code"] == "SESSION_EXPIRED"


@pytest.mark.parametrize("change", ["foreign_source", "foreign_fact", "duplicate_id", "closed", "resolution", "layout", "block"])
def test_preflight_issue_cannot_escape_selected_sources_or_claim_user_resolution(client, settings, monkeypatch, change):
    from app.models import Issue
    sid = _session(client)
    selected = _upload(client, sid, ("a.txt", SOURCE_A))
    rev = _select(client, sid, selected)
    original = MockAgent.analyze

    async def tampered(self, request):
        result = await original(self, request)
        issue = Issue(issue_id="tampered", scope="content", code="UNSUPPORTED_CLAIM", severity="blocker",
                      message="검사용 문제", source_ids=selected)
        if change == "foreign_source":
            issue.source_ids = ["source_other_session"]
        elif change == "foreign_fact":
            issue.fact_ids = ["fact_other_preflight"]
        elif change == "duplicate_id":
            issue.issue_id = result.issues[0].issue_id
        elif change == "closed":
            issue.status = "resolved"
        elif change == "resolution":
            issue.resolution = {"action": "acknowledged", "reason": "Agent 임의 확인"}
        elif change == "layout":
            issue.scope = "layout"
        elif change == "block":
            issue.block_ids = ["block_not_created"]
        result.issues.append(issue)
        return result

    monkeypatch.setattr(MockAgent, "analyze", tampered)
    response = client.post(f"/api/v1/sessions/{sid}/preflights", json={"expected_input_revision": rev})
    job = client.get(f"/api/v1/sessions/{sid}/jobs/{response.json()['job_id']}").json()
    assert job["status"] == "failed" and job["error"]["code"] == "AGENT_OUTPUT_INVALID"
    with connect(settings.db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM preflights WHERE session_id=?", (sid,)).fetchone()[0] == 0


def test_agent_output_with_unknown_segment_is_rejected(client, settings, monkeypatch):
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A))
    rev = _select(client, sid, src)
    original = MockAgent.analyze

    async def tampered(self, request):
        result = await original(self, request)
        fact = next(f for f in result.facts if f.status == "supported")
        fact.evidence_refs[0] = replace_ref(fact.evidence_refs[0], segment_id="seg_not_in_session")
        return result

    monkeypatch.setattr(MockAgent, "analyze", tampered)
    r = client.post(f"/api/v1/sessions/{sid}/preflights", json={"expected_input_revision": rev})
    job = client.get(f"/api/v1/sessions/{sid}/jobs/{r.json()['job_id']}").json()
    assert job["status"] == "failed" and job["error"]["code"] == "AGENT_OUTPUT_INVALID"


def replace_ref(ref: EvidenceRef, **kw) -> EvidenceRef:
    return EvidenceRef(**{**ref.model_dump(), **kw})


def test_agent_error_maps_to_job_error(client, monkeypatch):
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A))
    rev = _select(client, sid, src)

    async def boom(self, request):
        raise AgentError("AI_RATE_LIMIT", "한도 초과", retryable=True)

    monkeypatch.setattr(MockAgent, "analyze", boom)
    r = client.post(f"/api/v1/sessions/{sid}/preflights", json={"expected_input_revision": rev})
    job = client.get(f"/api/v1/sessions/{sid}/jobs/{r.json()['job_id']}").json()
    assert job["status"] == "failed" and job["error"] == {"code": "AI_RATE_LIMIT", "message": "한도 초과", "retryable": True,
                                                          "details": {}, "request_id": None}


def test_llm_mode_without_agent_impl_fails_not_mocks(tmp_path, monkeypatch):
    """실제 Agent 파일이 있어도 불러오기 실패를 재현하고 mock 대체를 막는다."""
    from unittest.mock import Mock

    import_stub = Mock()
    import_stub.import_module.side_effect = ModuleNotFoundError("테스트용 Agent 불러오기 실패")
    # 연결부의 importlib 참조만 교체한다. 다른 라이브러리의 import에는 영향을 주지 않는다.
    monkeypatch.setattr("app.agent_bridge.importlib", import_stub)
    settings = Settings(private_runs_dir=tmp_path / "runs", db_path=tmp_path / "runs" / "t.sqlite3", agent_mode="llm")
    client = TestClient(create_app(settings))
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A))
    rev = _select(client, sid, src)
    r = client.post(f"/api/v1/sessions/{sid}/preflights", json={"expected_input_revision": rev})
    job = client.get(f"/api/v1/sessions/{sid}/jobs/{r.json()['job_id']}").json()
    assert job["status"] == "failed" and job["error"]["code"] == "SERVICE_TEMPORARY_FAILURE"
    assert "agent_llm" in job["error"]["message"]
    import_stub.import_module.assert_called_once_with("app.agent_llm")


def test_sync_bridge_implementation_also_works(client, monkeypatch):
    """Agent가 동기 함수로 구현해도 붙는다."""
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A))
    rev = _select(client, sid, src)
    original = MockAgent.analyze

    def sync_analyze(self, request):
        import asyncio
        return asyncio.run(original(self, request))

    monkeypatch.setattr(MockAgent, "analyze", sync_analyze)
    assert _preflight(client, sid, rev)["can_generate"] is True


# ---------------- 초안 ----------------

def test_draft_requires_confirmation_and_matching_revision(client):
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A))
    rev = _select(client, sid, src)
    pf = _preflight(client, sid, rev)
    r = client.post(f"/api/v1/sessions/{sid}/drafts", json={"preflight_id": pf["preflight_id"], "input_revision": rev, "confirmed": False})
    assert r.status_code == 422 and r.json()["error"]["code"] == "PREFLIGHT_NOT_CONFIRMED"
    r = client.post(f"/api/v1/sessions/{sid}/drafts", json={"preflight_id": pf["preflight_id"], "input_revision": rev - 1, "confirmed": True})
    assert r.status_code == 409
    r = client.post(f"/api/v1/sessions/{sid}/drafts", json={"preflight_id": "pf_nope", "input_revision": rev, "confirmed": True})
    assert r.status_code == 404
    # 입력이 바뀌면 옛 점검으로 생성 불가(QA-07)
    _select(client, sid, src, expected=rev)
    r = client.post(f"/api/v1/sessions/{sid}/drafts", json={"preflight_id": pf["preflight_id"], "input_revision": rev + 1, "confirmed": True})
    assert r.status_code == 409 and r.json()["error"]["details"]["preflight_input_revision"] == rev


def test_draft_creates_document_rev1_with_evidence_and_asset(client, settings):
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A), ("photo.png", _png()))
    rev = _select(client, sid, src)
    pf = _preflight(client, sid, rev)
    r = client.post(f"/api/v1/sessions/{sid}/drafts", json={"preflight_id": pf["preflight_id"], "input_revision": rev, "confirmed": True})
    assert r.status_code == 202 and r.json()["kind"] == "draft"
    job = client.get(f"/api/v1/sessions/{sid}/jobs/{r.json()['job_id']}").json()
    assert job["status"] == "succeeded" and job["result_ref"]["type"] == "document"
    did = job["result_ref"]["document_id"]
    assert client.get(f"/api/v1/sessions/{sid}/preflights/{pf['preflight_id']}").json()["confirmed_at"] is not None
    out = client.get(f"/api/v1/sessions/{sid}/documents/{did}").json()
    doc = out["document"]
    assert out["validation"] is None and out["approval"] is None
    assert doc["schema_version"] == "1.0" and doc["document_revision"] == 1 and doc["input_revision"] == rev
    assert doc["title"] == "예시 회사" and doc["target_pages"] == 4 and len(doc["pages"]) == 4
    assert doc["status"] == "draft"  # warning만 있고 blocker 없음
    first = doc["pages"][0]["blocks"]
    assert first[0]["type"] == "heading" and first[0]["content"]["text"] == "예시 회사" and first[0]["fact_ids"]
    assert first[0]["evidence_refs"][0]["segment_id"].startswith("seg_")
    image = next(b for b in first if b["type"] == "image")
    asset_id = image["content"]["asset_id"]
    assert client.get(f"/api/v1/sessions/{sid}/assets/{asset_id}").status_code == 200
    # 근거 없는 문단 없음
    for page in doc["pages"]:
        for b in page["blocks"]:
            if b["type"] == "paragraph":
                assert b["fact_ids"] or b["content"]["text"] == "추가 확인 필요"
    # 세션 요약에 문서가 보인다
    assert client.get(f"/api/v1/sessions/{sid}").json()["document_summary"] == {"document_id": did, "document_revision": 1, "status": "draft", "demo": False}
    # 버전 구조: documents 머리 + document_revisions 1행
    with connect(settings.db_path) as conn:
        assert conn.execute("SELECT current_revision FROM documents WHERE document_id=?", (did,)).fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM document_revisions WHERE document_id=?", (did,)).fetchone()[0] == 1
    # 두 번째 생성은 거부(기존 편집 덮어쓰기 금지)
    r = client.post(f"/api/v1/sessions/{sid}/drafts", json={"preflight_id": pf["preflight_id"], "input_revision": rev, "confirmed": True})
    assert r.status_code == 409 and r.json()["error"]["code"] == "DOCUMENT_EXISTS"


@pytest.mark.parametrize("case", ["split", "unavailable", "input_changed", "bad_reference"])
def test_editorial_draft_layout_preparation_saves_once_and_rechecks_boundaries(client, settings, monkeypatch, case):
    from app.models import EditorialRecord, FactSelection, PageDesign
    from app.services import export_render as er
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A), ("photo.png", _png()))
    rev = _select(client, sid, src)
    pf = _preflight(client, sid, rev)
    original_draft = MockAgent.draft
    calls = []
    async def editorial(self, request):
        calls.append("draft")
        result = await original_draft(self, request)
        for p in result.pages:
            p.design = PageDesign()
        result.editorial = EditorialRecord(input_revision=rev, requested_pages=4, generated_pages=len(result.pages),
            page_count_reason="시험 구성", selections=[FactSelection(fact_id=f.fact_id,
                disposition="optional" if f.status == "supported" else "review", reason="시험 근거") for f in request.preflight.facts])
        return result
    def paginate(snapshot, out_dir, config):
        calls.append("paginate")
        assert out_dir.is_dir() and config.db_path == settings.db_path
        # Acquiring a second writer also verifies browser work holds no DB write lock.
        with connect(settings.db_path, immediate=True) as conn:
            if case == "input_changed":
                conn.execute("UPDATE sessions SET input_revision=input_revision+1 WHERE session_id=?", (sid,))
        pages = [p.model_copy(deep=True) for p in snapshot.pages]
        if case == "split":
            extra = pages[0].model_copy(deep=True)
            extra.page_id = "page_continuation"
            extra.blocks = pages[0].blocks[2:]
            pages[0].blocks = pages[0].blocks[:2]
            pages.insert(1, extra)
        if case == "bad_reference":
            pages[0].blocks[0].fact_ids = ["nonexistent_fact"]
        return er.DraftPagination(pages, "unavailable" if case == "unavailable" else "passed", 1)
    monkeypatch.setattr(MockAgent, "draft", editorial)
    monkeypatch.setattr(er, "paginate_draft", paginate)
    payload = {"preflight_id": pf["preflight_id"], "input_revision": rev, "confirmed": True}
    response = client.post(f"/api/v1/sessions/{sid}/drafts", json=payload)
    assert response.status_code == 202
    job = client.get(f"/api/v1/sessions/{sid}/jobs/{response.json()['job_id']}").json()
    assert calls == ["draft", "paginate"]
    assert not list((settings.private_runs_dir / sid / "artifacts").glob("tmp_*"))
    with connect(settings.db_path) as conn:
        assert conn.execute("SELECT COUNT(*) FROM layout_checks").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM approvals").fetchone()[0] == 0
        if case in {"input_changed", "bad_reference"}:
            assert job["status"] == "failed"
            assert conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 0
            assert job["error"]["code"] == ("INPUT_REVISION_CONFLICT" if case == "input_changed" else "AGENT_OUTPUT_INVALID")
            return
    assert job["status"] == "succeeded"
    document = client.get(f"/api/v1/sessions/{sid}/documents/{job['result_ref']['document_id']}").json()["document"]
    assert len(document["pages"]) == (5 if case == "split" else 4)
    assert document["editorial"]["generated_pages"] == len(document["pages"])
    assert document["target_pages"] == document["editorial"]["requested_pages"] == 4
    assert document["status"] == "draft"  # Public status comes from validation/approval, not the stored preparation hint.
    if case == "unavailable":
        assert "완료하지 못했습니다" in document["editorial"]["page_count_reason"]
        with connect(settings.db_path) as conn:
            assert conn.execute("SELECT status FROM document_revisions WHERE document_id=?", (document["document_id"],)).fetchone()[0] == "review_required"
    assert client.post(f"/api/v1/sessions/{sid}/drafts", json=payload).status_code == 409
    assert calls == ["draft", "paginate"]


def test_draft_with_blocker_is_review_required_and_no_photo_gives_placeholder(client):
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A), ("b.txt", SOURCE_B))
    rev = _select(client, sid, src)
    pf = _preflight(client, sid, rev)
    r = client.post(f"/api/v1/sessions/{sid}/drafts", json={"preflight_id": pf["preflight_id"], "input_revision": rev, "confirmed": True})
    job = client.get(f"/api/v1/sessions/{sid}/jobs/{r.json()['job_id']}").json()
    doc = client.get(f"/api/v1/sessions/{sid}/documents/{job['result_ref']['document_id']}").json()["document"]
    assert doc["status"] == "review_required" and doc["title"] == "예시 회사" or doc["title"] == "예시 회사"
    assert any(b["type"] == "image_placeholder" for b in doc["pages"][0]["blocks"])


def test_draft_output_with_foreign_asset_is_rejected(client, monkeypatch):
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A))
    rev = _select(client, sid, src)
    pf = _preflight(client, sid, rev)
    original = MockAgent.draft

    async def tampered(self, request):
        result = await original(self, request)
        result.pages[0].blocks.append(Block(block_id="block_x", type="image",
                                            content={"asset_id": "asset_other", "alt": "", "caption": "", "fit": "contain"}))
        return result

    monkeypatch.setattr(MockAgent, "draft", tampered)
    r = client.post(f"/api/v1/sessions/{sid}/drafts", json={"preflight_id": pf["preflight_id"], "input_revision": rev, "confirmed": True})
    job = client.get(f"/api/v1/sessions/{sid}/jobs/{r.json()['job_id']}").json()
    assert job["status"] == "failed" and job["error"]["code"] == "AGENT_OUTPUT_INVALID"
    assert client.get(f"/api/v1/sessions/{sid}").json()["document_summary"] is None


def test_other_owner_cannot_read_preflight_or_document(client, settings):
    sid = _session(client)
    src = _upload(client, sid, ("a.txt", SOURCE_A))
    rev = _select(client, sid, src)
    pf = _preflight(client, sid, rev)
    r = client.post(f"/api/v1/sessions/{sid}/drafts", json={"preflight_id": pf["preflight_id"], "input_revision": rev, "confirmed": True})
    did = client.get(f"/api/v1/sessions/{sid}/jobs/{r.json()['job_id']}").json()["result_ref"]["document_id"]
    other = TestClient(create_app(settings))
    other.post("/api/v1/sessions", json={"brief": BRIEF})
    assert other.get(f"/api/v1/sessions/{sid}/preflights/{pf['preflight_id']}").status_code == 404
    assert other.get(f"/api/v1/sessions/{sid}/documents/{did}").status_code == 404
    assert other.post(f"/api/v1/sessions/{sid}/preflights", json={"expected_input_revision": rev}).status_code == 404


def test_mock_never_contains_real_company_terms():
    """mock 고정 문구에 실제 회사 정보가 없는지(값은 전부 입력 구간에서만 온다)."""
    import inspect
    import app.agent_mock as m
    text = inspect.getsource(m)
    for banned in ("거산", "케미칼", "Geosan"):
        assert banned not in text


def test_sufficiency_uses_evidence_and_preserves_conflicts():
    from app.models import Fact, Issue, EvidenceRef
    from app.services.preflights import assess_sufficiency
    ref = EvidenceRef(source_id="s", source_version=1, segment_id="seg", locator={}, excerpt="근거")
    def fact(key, status="supported", evidence=True):
        return Fact(fact_id=key, field_key=key, value="자료", status=status, evidence_refs=[ref] if evidence else [])
    facts = [fact("company_name"), fact("processes"), fact("customers_markets"), fact("certifications")]
    assert assess_sufficiency(facts, []).score == 100
    facts[3] = fact("certifications", evidence=False)
    assert assess_sufficiency(facts, []).score == 75
    facts[1] = fact("processes", status="conflict")
    facts.append(fact("technology"))
    result = assess_sufficiency(facts, [])
    assert result.score == 50 and result.categories[1].status == "conflict"
    issue = Issue(issue_id="i", scope="content", code="UNSUPPORTED_CLAIM", severity="blocker",
                  message="확인", fact_ids=["company_name"])
    result = assess_sufficiency(facts, [issue])
    assert result.score == 25 and result.has_blockers
    assert result.categories[0].status == "needs_confirmation"
    issue.scope = "source"
    issue.source_ids = ["s"]
    assert assess_sufficiency(facts, [issue]).score == 0
    assert assess_sufficiency([], []).score == 0


def test_preflight_api_exposes_evidence_coverage(client):
    sid = _session(client)
    ids = _upload(client, sid, ("company.txt", SOURCE_A))
    rev = _select(client, sid, ids)
    result = _preflight(client, sid, rev)
    coverage = result["sufficiency"]
    assert len(coverage["categories"]) == 4
    assert coverage["score"] == 25 * sum(c["status"] == "supported" for c in coverage["categories"])
    assert isinstance(coverage["has_blockers"], bool)


@pytest.mark.parametrize("case,ready", [
    ("same", True), ("corporation", True), ("other", False),
    ("unknown_translation", False), ("missing_name", False),
    ("conflict", False), ("missing_required_excluded", False),
    ("missing_required", True), ("excluded_opening", True),
    ("empty_supported", False),
])
def test_draft_entry_preflight_and_post_use_same_conditions(client, settings, monkeypatch, case, ready):
    from app.services import preflights
    sid = _session(client)
    ids = _upload(client, sid, ("company.txt", SOURCE_A))
    rev = _select(client, sid, ids)
    pf = _preflight(client, sid, rev)
    brief = dict(BRIEF, target_company="예시 회사")
    if case == "corporation": brief["target_company"] = "㈜예시 회사"
    if case == "other": brief["target_company"] = "다른회사"
    if case == "unknown_translation": brief["target_company"] = "EXAMPLE COMPANY"
    if case in {"conflict", "missing_required_excluded", "missing_required"}:
        brief["required_fields"] = ["certifications"]
        if case != "missing_required": brief["emphasis"] = ["인증 제외"]
    if case == "excluded_opening": brief["emphasis"] = ["회사 개요 제외"]
    with connect(settings.db_path) as conn:
        conn.execute("UPDATE sessions SET brief_json=? WHERE session_id=?", (json.dumps(brief), sid))
        facts = pf["facts"]
        for fact in facts:
            if (case == "missing_name" and fact["field_key"] == "company_name") or case == "empty_supported":
                fact.update(status="missing", value=None, evidence_refs=[])
        conn.execute("UPDATE preflights SET facts_json=? WHERE preflight_id=?", (json.dumps(facts),pf["preflight_id"]))
        stored = preflights.get(conn, sid, pf["preflight_id"]).model_dump()
        job_count = conn.execute("SELECT count(*) FROM jobs WHERE session_id=?", (sid,)).fetchone()[0]
    client.app.state.settings = replace(settings, agent_mode="llm")
    monkeypatch.setattr(ai_jobs, "get_bridge", lambda s: MockAgent())
    observed = client.get(f"/api/v1/sessions/{sid}/preflights/{pf['preflight_id']}").json()
    assert observed["can_generate"] is ready
    if not ready:
        assert observed["recommendations"]["needed"]
    with connect(settings.db_path) as conn:
        assert preflights.get(conn,sid,pf["preflight_id"]).model_dump() == stored  # GET never rewrites facts/confirmation.
    r = client.post(f"/api/v1/sessions/{sid}/drafts",json={"preflight_id":pf["preflight_id"],"input_revision":rev,"confirmed":True})
    if ready:
        assert r.status_code == 202, r.text
    else:
        assert r.status_code == 422, r.text
        assert r.json()["error"]["details"]["recovery_action"] == "review_inputs"
        with connect(settings.db_path) as conn:
            assert conn.execute("SELECT count(*) FROM jobs WHERE session_id=?", (sid,)).fetchone()[0] == job_count
            assert preflights.get(conn,sid,pf["preflight_id"]).confirmed_at is None
            assert not conn.execute("SELECT 1 FROM documents WHERE session_id=?",(sid,)).fetchone()


def test_older_same_input_preflight_is_not_offered_or_consumed(client, settings):
    from app.services import preflights
    sid = _session(client)
    rev = _select(client,sid,_upload(client,sid,("company.txt",SOURCE_A)))
    first = _preflight(client,sid,rev)
    latest = _preflight(client,sid,rev)
    view = client.get(f"/api/v1/sessions/{sid}/preflights/{first['preflight_id']}").json()
    assert not view["can_generate"] and view["recommendations"]["needed"]
    assert view["latest_preflight_id"] == latest["preflight_id"]
    assert client.get(f"/api/v1/sessions/{sid}/preflights/{latest['preflight_id']}").json()["can_generate"]
    with connect(settings.db_path) as conn:
        count = conn.execute("SELECT count(*) FROM jobs WHERE session_id=?",(sid,)).fetchone()[0]
    response = client.post(f"/api/v1/sessions/{sid}/drafts",json={"preflight_id":first['preflight_id'],"input_revision":rev,"confirmed":True})
    assert response.status_code == 409
    assert response.json()["error"]["details"]["latest_preflight_id"] == latest['preflight_id']
    with connect(settings.db_path) as conn:
        assert preflights.get(conn,sid,first['preflight_id']).confirmed_at is None
        assert conn.execute("SELECT count(*) FROM jobs WHERE session_id=?",(sid,)).fetchone()[0] == count


@pytest.mark.parametrize("orm", [False, True])
def test_preflight_review_exclude_restore_replay_and_required_guards(settings, orm):
    if orm:
        init_orm_db(settings.db_path, settings.private_runs_dir)
    with TestClient(create_app(settings)) as client:
        sid = _session(client)
        src = _upload(client, sid, ("a.txt", SOURCE_A))
        rev = _select(client, sid, src)
        original = _preflight(client, sid, rev)
        base = f"/api/v1/sessions/{sid}/preflights/"
        optional = next(fact for fact in original["facts"] if fact["field_key"] == "process_count")
        body = {"expected_input_revision": rev, "action": "exclude", "fact_ids": [optional["fact_id"]], "reason": "이번 문서에서 사용하지 않음"}
        url = base + original["preflight_id"] + "/reviews"
        headers = {"Idempotency-Key": "review-exclude"}
        response = client.post(url, json=body, headers=headers)
        assert response.status_code == 200, response.text
        reviewed = response.json()
        assert reviewed["preflight_id"] != original["preflight_id"] and reviewed["confirmed_at"] is None
        assert optional not in reviewed["facts"] and reviewed["excluded_facts"] == [optional]
        assert client.post(url, json=body, headers=headers).json() == reviewed
        assert client.get(base + original["preflight_id"]).json()["facts"] == original["facts"]
        stale = client.post(url, json=body)
        assert stale.status_code == 409
        latest_url = base + reviewed["preflight_id"] + "/reviews"
        company = next(fact for fact in reviewed["facts"] if fact["field_key"] == "company_name")
        rejected = client.post(latest_url, json={**body, "fact_ids": [company["fact_id"]]})
        assert rejected.status_code == 422 and rejected.json()["error"]["code"] == "RESOLUTION_NOT_ALLOWED"
        businesses = [fact["fact_id"] for fact in reviewed["facts"] if fact["status"] == "supported" and fact["field_key"] in {"company_summary", "business_areas", "products_services", "technology", "processes"}]
        assert client.post(latest_url, json={**body, "fact_ids": businesses}).status_code == 422
        assert client.post(latest_url, json={**body, "fact_ids": ["foreign_fact"]}).status_code == 422
        restored = client.post(latest_url, json={**body, "action": "restore"})
        assert restored.status_code == 200, restored.text
        assert restored.json()["excluded_facts"] == [] and restored.json()["facts"] == original["facts"]
        other = TestClient(create_app(settings))
        other.post("/api/v1/sessions", json={"brief": BRIEF})
        assert other.post(base + restored.json()["preflight_id"] + "/reviews", json=body).status_code == 404


@pytest.mark.parametrize("orm", [False, True])
def test_excluded_preflight_fact_is_absent_from_draft_and_invalid_when_reintroduced(settings, orm):
    from app.models import Document
    from app.services import preflights, validation
    if orm:
        init_orm_db(settings.db_path, settings.private_runs_dir)
    with TestClient(create_app(settings)) as client:
        sid = _session(client); src = _upload(client, sid, ("a.txt", SOURCE_A)); rev = _select(client, sid, src)
        pf = _preflight(client, sid, rev)
        optional = next(fact for fact in pf["facts"] if fact["field_key"] == "process_count")
        base = f"/api/v1/sessions/{sid}"
        reviewed = client.post(base + f"/preflights/{pf['preflight_id']}/reviews", json={
            "expected_input_revision": rev, "action": "exclude", "fact_ids": [optional["fact_id"]], "reason": "제외"}).json()
        accepted = client.post(base + "/drafts", json={"input_revision": rev, "preflight_id": reviewed["preflight_id"], "confirmed": True})
        assert accepted.status_code == 202, accepted.text
        job = client.get(base + "/jobs/" + accepted.json()["job_id"]).json(); assert job["status"] == "succeeded", job
        doc = client.get(base + "/documents/" + job["result_ref"]["document_id"]).json()["document"]
        assert all(optional["fact_id"] not in block["fact_ids"] for page in doc["pages"] for block in page["blocks"])
        with connect(settings.db_path) as conn:
            result = preflights.get(conn, sid, reviewed["preflight_id"])
            ctx = validation.load_context(conn, sid, result)
            document = Document.model_validate(doc)
            block = document.pages[0].blocks[0]
            block.fact_ids = [optional["fact_id"]]
            block.evidence_refs = [EvidenceRef.model_validate(ref) for ref in optional["evidence_refs"]]
            findings, _ = validation.server_checks(document, ctx)
            assert any(issue.code == "EVIDENCE_INVALID" and block.block_id in issue.block_ids for issue in findings)


@pytest.mark.parametrize("orm", [False, True])
def test_optional_conflict_exclusion_closes_only_related_issue_and_restore_reopens(settings, orm):
    from app.services import preflights, validation, jobs
    if orm:
        init_orm_db(settings.db_path, settings.private_runs_dir)
    with TestClient(create_app(settings)) as client:
        sid = _session(client)
        src = _upload(client, sid, ("a.txt", SOURCE_A), ("b.txt", "회사명: 예시 회사\n납기 표현: 느린 납기\n".encode()))
        rev = _select(client, sid, src)
        pf = _preflight(client, sid, rev)
        fact = next(fact for fact in pf["facts"] if fact["field_key"] == "lead_time")
        assert fact["status"] == "conflict"
        base = f"/api/v1/sessions/{sid}/preflights/"
        body = {"expected_input_revision": rev, "action": "exclude", "fact_ids": [fact["fact_id"]], "reason": "선택 항목 제외"}
        changed = client.post(base + pf["preflight_id"] + "/reviews", json=body)
        assert changed.status_code == 200, changed.text
        changed = changed.json()
        assert not any(issue["status"] == "open" and fact["fact_id"] in issue["fact_ids"] for issue in changed["issues"])
        assert changed["excluded_facts"][0]["status"] == "conflict"
        with connect(settings.db_path) as conn:
            ctx = validation.load_context(conn, sid, preflights.get(conn, sid, changed["preflight_id"]))
            assert not any(fact["fact_id"] in issue.fact_ids for issue in validation.preflight_conflicts(ctx))
        restored = client.post(base + changed["preflight_id"] + "/reviews", json={**body, "action": "restore"}).json()
        assert any(issue["status"] == "open" and fact["fact_id"] in issue["fact_ids"] for issue in restored["issues"])
        with connect(settings.db_path, immediate=True) as conn:
            active = jobs.create(conn, sid, "draft", "진행 중", input_revision=rev)
        rejected = client.post(base + restored["preflight_id"] + "/reviews", json=body)
        assert rejected.status_code == 409
        with connect(settings.db_path, immediate=True) as conn:
            conn.execute("UPDATE jobs SET status='cancelled' WHERE job_id=?", (active.job_id,))
        changed_rev = client.patch(f"/api/v1/sessions/{sid}/inputs", json={"expected_input_revision": rev,
            "brief": {**BRIEF, "required_fields": ["lead_time"]}}).json()["input_revision"]
        assert client.post(base + restored["preflight_id"] + "/reviews", json=body).status_code == 409
        latest = _preflight(client, sid, changed_rev)
        fact = next(fact for fact in latest["facts"] if fact["field_key"] == "lead_time")
        assert client.post(base + latest["preflight_id"] + "/reviews", json={**body,
            "expected_input_revision": changed_rev, "fact_ids": [fact["fact_id"]]}).status_code == 422

@pytest.mark.parametrize("orm", [False, True])
@pytest.mark.parametrize("configured_mode", ["mock", "llm"])
def test_job_trace_preserves_input_metadata_counts_executor_and_no_source_text(settings, monkeypatch, orm, configured_mode):
    from app.services import preflights
    settings.private_runs_dir.mkdir(parents=True, exist_ok=True)
    if orm:
        init_orm_db(settings.db_path, settings.private_runs_dir)
    settings = replace(settings, agent_mode=configured_mode)
    seen = []
    original = MockAgent.analyze
    async def capture(self, request):
        seen.append(request)
        return await original(self, request)
    monkeypatch.setattr(MockAgent, "analyze", capture)
    monkeypatch.setattr(ai_jobs, "get_bridge", lambda _: MockAgent())
    with TestClient(create_app(settings)) as api:
        sid = _session(api)
        text = "회사명: 예시 회사\n회사 개요: 제조 시험 기업\n사업 분야: 표면처리\n시험 값: 168시간"
        src = _upload(api, sid, ("secret-original-name.txt", text.encode()))
        with connect(settings.db_path) as conn:
            conn.execute("UPDATE sources SET document_date='2017', document_date_verified=0, date_from_filename='2025' WHERE source_id=?", (src[0],))
            conn.execute("UPDATE segments SET evidence_status='unverified', document_date='2016', chunk_id='chunk_fixed' WHERE source_id=?", (src[0],))
        revision = _select(api, sid, src)
        r = api.post(f"/api/v1/sessions/{sid}/preflights", json={"expected_input_revision": revision})
        job = api.get(f"/api/v1/sessions/{sid}/jobs/{r.json()['job_id']}").json()
        assert job["status"] == "succeeded", job
        trace = job["progress"]["trace"]
        stages = {s["stage"]: s for s in trace["stages"]}
        assert trace["input_revision"] == revision and trace["outcome"]["status"] == "succeeded"
        assert stages["executor"]["configured_mode"] == configured_mode
        assert stages["executor"]["actual_mode"] == "mock" and stages["executor"]["model"] is None
        assert stages["agent_input"]["source_count"] == 1
        source = seen[0].sources[0]
        assert source.source_id == src[0] and source.source_version == 1
        assert stages["agent_input"]["segment_count"] == len(source.segments)
        assert stages["agent_input"]["input_text_chars"] == sum(len(s.text) for s in source.segments)
        assert source.metadata["document_date"] == "2017" and source.metadata["document_date_verified"] is False
        assert source.metadata["date_from_filename"] == "2025"
        assert all(m["document_date"] == "2016" and m["evidence_status"] == "unverified"
                   and m["chunk_id"] == "chunk_fixed" for m in source.metadata["segments"].values())
        serialized = json.dumps(trace, ensure_ascii=False)
        assert "secret-original-name" not in serialized and "예시 회사" not in serialized and "168시간" not in serialized
        assert "unverified" not in serialized  # Status presence counts only, no arbitrary source-authored labels.
        assert stages["result_checked"]["fact_status_counts"]["supported"] >= 2
        # Missing/unavailable selections are explained without exposing the other source.
        with connect(settings.db_path) as conn:
            excluded = []
            assert preflights.build_sources(conn, sid, ["foreign_source"], exclusions=excluded) == []
            assert excluded == [{"source_id": "foreign_source", "reason": "outside_evidence_scope"}]
            if orm:
                conn.execute("UPDATE sources SET source_version=2, current_run_id=NULL, document_date='2099', date_from_filename='2099' WHERE source_id=?", (src[0],))
                pinned = preflights.build_sources(conn, sid, src)[0]
                assert pinned.source_version == 1 and pinned.metadata["document_date"] is None
                assert pinned.metadata["document_date_verified"] is None and pinned.metadata["date_from_filename"] is None
                assert all(m["document_date"] == "2016" for m in pinned.metadata["segments"].values())


@pytest.mark.parametrize("outcome", ["failed", "cancelled", "succeeded"])
def test_job_trace_terminal_guards_and_session_purge(client, settings, outcome):
    from app.services import jobs
    sid = _session(client)
    with connect(settings.db_path) as conn:
        job = jobs.create(conn, sid, "preflight", "pending", input_revision=1)
        assert jobs.record_trace(conn, job.job_id, "agent_input", {"source_count": 2}) == 1
        jobs.set_progress(conn, job.job_id, "analyzing", "working")
        if outcome == "failed": jobs.fail(conn, job.job_id, "AGENT_OUTPUT_INVALID", "fixed", True)
        elif outcome == "cancelled": jobs.cancel_for_session(conn, sid)
        else: jobs.succeed(conn, job.job_id, {"type": "preflight", "preflight_id": "pf_test"})
        before = conn.execute("SELECT progress_json FROM jobs WHERE job_id=?", (job.job_id,)).fetchone()[0]
        assert jobs.record_trace(conn, job.job_id, "late", {"source_count": 99}) == 0
        assert jobs.set_progress(conn, job.job_id, "late", None) == 0
        assert jobs.succeed(conn, job.job_id, {}) == 0 and jobs.fail(conn, job.job_id, "LATE", "fixed", True) == 0
        assert conn.execute("SELECT progress_json FROM jobs WHERE job_id=?", (job.job_id,)).fetchone()[0] == before
        assert jobs.get(conn, sid, job.job_id).progress.trace is not None
        jobs.purge_errors_for_session(conn, sid)
        assert jobs.get(conn, sid, job.job_id).progress.trace is None
        assert jobs.get(conn, sid, job.job_id).status == outcome


def test_job_trace_bound_is_explicit_not_silent(client, settings):
    from app.services import jobs
    sid = _session(client)
    with connect(settings.db_path) as conn:
        job = jobs.create(conn, sid, "read", "pending")
        for i in range(20): jobs.record_trace(conn, job.job_id, f"parsed:{i}", {"segment_count": i})
        trace = jobs.get(conn, sid, job.job_id).progress.trace
        assert len(trace["stages"]) == 16 and trace["omitted_stage_count"] == 4
        jobs.record_trace(conn, job.job_id, "agent_input", {"source_id": "x" * 70000})
        trace = jobs.get(conn, sid, job.job_id).progress.trace
        assert trace["stages"][0]["summary_omitted"] == "trace_size_limit"

@pytest.mark.parametrize("purpose", ["품질 담당자용 인증과 시험 근거", "생산기술 담당자용 공정과 설비 검토", "신규 고객 소개용 주요 서비스"])
def test_customer_request_table_transfer_and_job_trace_without_llm_quality_claim(client, settings, purpose):
    from test_be03 import make_docx
    from app.services import preflights
    table = [["공정/설비/시험", "대상과 조건"], ["아노다이징", "알루미늄 시편"],
             ["시험 장비 A", "시편 2개 조건"], ["시험 시간", "168시간; 시험 종류 미기재"]]
    data = make_docx(["회사명: 예시 회사", "회사 개요: 제조업 시험 기업", "사업 분야: 표면처리"], table)
    sid = client.post("/api/v1/sessions", json={"brief": BRIEF | {"purpose": purpose, "target_pages": 1}}).json()["session_id"]
    src = _upload(client, sid, ("manufacturing-fixture.docx", data))
    revision = _select(client, sid, src)
    pf = _preflight(client, sid, revision)
    with connect(settings.db_path) as conn:
        source = preflights.build_sources(conn, sid, src)[0]
    cells = [g for g in source.segments if "table" in g.locator]
    assert len(cells) == 8
    assert {(g.locator["row"], g.locator["col"]): g.text for g in cells} == {
        (r, c): value for r, row in enumerate(table, 1) for c, value in enumerate(row, 1)}
    assert all(source.metadata["segments"][g.segment_id]["extraction_method"] == "parser" for g in source.segments)
    accepted = client.post(f"/api/v1/sessions/{sid}/drafts", json={"preflight_id": pf["preflight_id"],
        "input_revision": revision, "confirmed": True})
    assert accepted.status_code == 202, accepted.text
    job = client.get(f"/api/v1/sessions/{sid}/jobs/{accepted.json()['job_id']}").json()
    assert job["status"] == "succeeded", job
    stages = {s["stage"]: s for s in job["progress"]["trace"]["stages"]}
    assert stages["agent_input"]["preflight_id"] == pf["preflight_id"]
    assert stages["agent_input"]["segment_count"] == 11
    assert stages["executor"]["actual_mode"] == "mock" and stages["executor"]["model"] is None
    assert stages["result_checked"]["page_count"] >= 1
    assert job["result_ref"]["document_revision"] == 1
    document = client.get(f"/api/v1/sessions/{sid}/documents/{job['result_ref']['document_id']}").json()["document"]
    for page in document["pages"]:
        for block in page["blocks"]:
            assert all(r["source_id"] == src[0] and r["source_version"] == 1 for r in block["evidence_refs"])


@pytest.mark.parametrize("case_id", CUSTOMER_QUALITY_CASES)
def test_customer_quality_fixture_preserves_conditions_without_sending_rubric(client, settings, monkeypatch, case_id):
    from app.services import preflights
    captured = []
    original = MockAgent.analyze

    def observe(self, request):
        captured.append(request)
        return original(self, request)

    monkeypatch.setattr(MockAgent, "analyze", observe)
    brief = customer_quality_brief(case_id)
    sid = client.post("/api/v1/sessions", json={"brief": brief}).json()["session_id"]
    ids = _upload(client, sid, ("synthetic-manufacturing.txt", CUSTOMER_QUALITY_SOURCE.encode()), kind="company")
    revision = _select(client, sid, ids)
    _preflight(client, sid, revision)
    assert len(captured) == 1
    request = captured[0]
    assert request.brief.model_dump() == client.get(f"/api/v1/sessions/{sid}").json()["brief"]
    assert request.brief.required_fields == brief["required_fields"]
    with connect(settings.db_path) as conn:
        source = preflights.build_sources(conn, sid, ids)[0]
    assert request.sources == [source]
    parsed = "\n".join(segment.text for segment in source.segments)
    assert all(line in parsed for line in CUSTOMER_QUALITY_SOURCE.splitlines())
    # Probe IDs/expected presence are evaluation metadata, not company evidence.
    assert not any(key in source.metadata for key in ("probes", "expected", "rubric"))
    assert source.source_version == 1
    assert all(segment.segment_id and segment.locator for segment in source.segments)


@pytest.mark.parametrize("fmt", ["txt", "docx"])
def test_long_table_quality_input_keeps_core_conditions(client, monkeypatch, fmt):
    captured = []
    original = MockAgent.analyze

    def observe(self, request):
        captured.append(request)
        return original(self, request)

    monkeypatch.setattr(MockAgent, "analyze", observe)
    brief = {**BRIEF, "target_company": "예시유체", "photo_preference": "none"}
    sid = client.post("/api/v1/sessions", json={"brief": brief}).json()["session_id"]
    ids = _upload(client, sid, ("held-out." + fmt, long_table_quality_source(fmt)), kind="company")
    _preflight(client, sid, _select(client, sid, ids))
    assert len(captured) == 1
    source = captured[0].sources[0]
    parsed = "\n".join(segment.text for segment in source.segments)
    for key, value in LONG_TABLE_QUALITY_ROWS:
        assert key in parsed and value in parsed
    assert not any(key in source.metadata for key in ("probes", "expected", "rubric"))
    assert all(segment.segment_id and segment.locator for segment in source.segments)
    if fmt == "docx":
        assert len(parsed) > 4000
        assert parsed.count("자료 배경:") == 72


@pytest.mark.parametrize("case", MULTISOURCE_TIME_CASES)
def test_multisource_time_input_keeps_separate_dates_and_sources(client, monkeypatch, case):
    captured = []
    original = MockAgent.analyze

    def observe(self, request):
        captured.append(request)
        return original(self, request)

    monkeypatch.setattr(MockAgent, "analyze", observe)
    brief = dict(BRIEF, target_company="예시표면", required_fields=["capabilities"])
    sid = client.post("/api/v1/sessions", json={"brief": brief}).json()["session_id"]
    originals = MULTISOURCE_TIME_CASES[case]
    ids = _upload(client, sid, *((f"source-{index}.txt", text.encode()) for index, text in enumerate(originals)))
    _preflight(client, sid, _select(client, sid, ids))
    assert len(captured) == 1
    assert captured[0].brief.required_fields == ["capabilities"]
    assert {source.source_id for source in captured[0].sources} == set(ids)
    expected = dict(zip(ids, originals))
    for source in captured[0].sources:
        original_text = expected[source.source_id]
        parsed = "\n".join(segment.text for segment in source.segments)
        assert all(line in parsed for line in original_text.splitlines())
        assert source.source_version == 1
        assert all(segment.segment_id and segment.locator for segment in source.segments)


def _layout_pptx(*, duplicate=False, rotated=False, offset=0):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Deliberately put the bodies before the titles in XML order.
    entries = [(4, 2, "알루미늄 소재 후처리 내식성 168hr"),
               (4, 4, "금속 산화피막 설명"),
               (1, 2, "크로메이트"), (1, 4, "부동태")]
    if duplicate == True:
        entries.append((1, 6, "알루미늄 소재 후처리 내식성 168hr"))
    for x, y, text in entries:
        box = slide.shapes.add_textbox(Inches(x), Inches(y + offset), Inches(2), Inches(.5))
        box.text = text
        if rotated and text.startswith("알루미늄"):
            box.rotation = 90
    if duplicate in ("group", "rotated"):
        owner = slide.shapes.add_group_shape().shapes if duplicate == "group" else slide.shapes
        box = owner.add_textbox(Inches(1), Inches(6), Inches(2), Inches(.5))
        box.text = "알루미늄 소재 후처리 내식성 168hr"
        if duplicate == "rotated":
            box.rotation = 90
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("case", ["valid", "changed_text", "wrong_slide", "ambiguous", "rotated",
                                   "hash_mismatch", "missing_hash", "missing_file", "oversize", "outside_root", "corrupt", "group_duplicate", "rotated_duplicate"])
def test_pptx_context_never_attaches_unmatched_or_untrusted_geometry(tmp_path, case):
    import hashlib
    from app.agent_bridge import SegmentIn
    from app.services.preflights import _source_layout
    root = tmp_path / "runs"
    root.mkdir()
    settings = Settings(private_runs_dir=root, db_path=root / "test.sqlite3")
    data = b"not a pptx" if case == "corrupt" else _layout_pptx(duplicate=case.removesuffix("_duplicate") if case.endswith("_duplicate") else case == "ambiguous", rotated=case == "rotated")
    filename = "../outside.pptx" if case == "outside_root" else "deck.pptx"
    path = root / filename
    if case != "missing_file":
        path.write_bytes(data)
    digest = hashlib.sha256(data).hexdigest()
    if case == "hash_mismatch":
        digest = "0" * 64
    if case == "missing_hash":
        digest = None
    if case == "oversize":
        settings = replace(settings, max_file_bytes=len(data) - 1)
    segment = SegmentIn("selected", {"slide": 2 if case == "wrong_slide" else 1},
                        "없는 설명" if case == "changed_text" else "알루미늄 소재 후처리 내식성 168hr")
    result = _source_layout(settings, filename, digest, [segment])
    if case == "valid":
        assert result == {"selected": {"slide": 1, "box_mm": [101.6, 50.8, 50.8, 12.7]}}
        assert "크로메이트" not in json.dumps(result, ensure_ascii=False)  # unselected title is not leaked
    else:
        assert result == {}
    assert segment.locator == {"slide": 2 if case == "wrong_slide" else 1}


@pytest.mark.parametrize("orm", [False, True])
def test_preflight_passes_pptx_layout_from_selected_version_without_rewriting_evidence(settings, monkeypatch, orm):
    import hashlib
    from app.services import preflights
    if orm:
        settings.private_runs_dir.mkdir(parents=True, exist_ok=True)
        init_orm_db(settings.db_path, settings.private_runs_dir)
    seen = []
    original = MockAgent.analyze
    async def capture(self, request):
        seen.append(request)
        return await original(self, request)
    monkeypatch.setattr(MockAgent, "analyze", capture)
    monkeypatch.setattr(ai_jobs, "get_bridge", lambda _: MockAgent())
    with TestClient(create_app(settings)) as api:
        sid = _session(api)
        ids = _upload(api, sid, ("two-processes.pptx", _layout_pptx()))
        rev = _select(api, sid, ids)
        with connect(settings.db_path) as conn:
            before = preflights.build_sources(conn, sid, ids)[0]
            # Existing registered excerpts may have only a slide locator.
            for segment in before.segments:
                conn.execute("UPDATE segments SET locator_json=? WHERE segment_id=?", ('{"slide":1}', segment.segment_id))
            before = preflights.build_sources(conn, sid, ids)[0]
            if orm:
                data = _layout_pptx(offset=1)
                (settings.private_runs_dir / "new-version.pptx").write_bytes(data)
                conn.execute("UPDATE sources SET source_version=2, current_run_id=NULL, stored_path=?, content_hash=? WHERE source_id=?",
                             ("new-version.pptx", hashlib.sha256(data).hexdigest(), ids[0]))
        _preflight(api, sid, rev)
        source = seen[0].sources[0]
        assert source.source_version == 1 and source.segments == before.segments
        title = next(s for s in source.segments if s.text == "크로메이트")
        body = next(s for s in source.segments if "168hr" in s.text)
        assert source.metadata["segments"][title.segment_id]["layout"]["box_mm"] == [25.4, 50.8, 50.8, 12.7]
        assert source.metadata["segments"][body.segment_id]["layout"]["box_mm"] == [101.6, 50.8, 50.8, 12.7]
        with connect(settings.db_path) as conn:
            after = preflights.build_sources(conn, sid, ids)[0]
            assert after.segments == before.segments
            # A foreign session's selected IDs cannot bring their file context in.
            other = _session(api)
            assert preflights.build_sources(conn, other, ids, settings=settings) == []


# Held-out layout cases: expected relationships are test data, never prompt input.
def _pptx_layout_quality_case(case):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(12), Inches(7)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = "기업명: 예시배치기업. 사업: 산업용 표면처리 제품 제조."
    body_a = "알루미늄 하우징용. 시험 조건: 22도, 48시간. 출하 성능 보증값 아님."
    body_b = "스테인리스 부품용. 검사 조건: 35도, 12시간. 출하 성능 보증값 아님."
    event = "가공설비 사진. 2018년에 제2공장 준공. 2022년 설비 교체."
    if case in ("columns", "mirrored_reordered"):
        ax, bx = (1, 6) if case == "columns" else (6, 1)
        entries = [(bx, 2, body_b), (ax, 1, "막처리 A"), (ax, 2, body_a), (bx, 1, "세정 B"), (1, 4, event)]
        if case == "mirrored_reordered":
            entries.reverse()
        expected = {"a": body_a, "b": body_b, "event": event}
    elif case == "ambiguous":
        # Overlapping titles alone cannot determine which process the note describes.
        note = "공정 적용 소재는 알루미늄 또는 스테인리스 중 확인 필요. 검사시간은 12시간 또는 48시간으로 기록되었으나 공정별 대응은 미확정."
        entries = [(1, 1, "막처리 A"), (1, 1, "세정 B"), (1, 2, note)]
        expected = {"ambiguous": note}
    elif case == "conflict":
        one = "설비대장: 2024년 12월 31일 A공장 가공설비 4대."
        two = "현장조사: 2024년 12월 31일 A공장 가공설비 6대."
        entries = [(1, 1, one), (1, 3, two)]
        expected = {"conflict": [one, two]}
    else:
        raise ValueError(case)
    entries.append((1, .1, company))
    for x, y, text in entries:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(4), Inches(.7))
        box.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), expected


@pytest.mark.parametrize("case", ["columns", "mirrored_reordered", "ambiguous", "conflict"])
def test_pptx_layout_varied_cases_preserve_original_context_without_invented_relations(tmp_path, case):
    import hashlib
    from app.agent_bridge import SegmentIn
    from app.parsers import parse
    from app.services.preflights import _source_layout
    data, expected = _pptx_layout_quality_case(case)
    path = tmp_path / "fixture.pptx"
    path.write_bytes(data)
    settings = Settings(private_runs_dir=tmp_path, db_path=tmp_path / "unused.sqlite3")
    parsed = parse(data, ".pptx")
    assert parsed.status == "complete"
    segments = [SegmentIn(f"s{i}", s.locator.copy(), s.text) for i, s in enumerate(parsed.segments)]
    before = [(s.segment_id, s.locator.copy(), s.text) for s in segments]
    layout = _source_layout(settings, path.name, hashlib.sha256(data).hexdigest(), segments)
    assert set(layout) == {s.segment_id for s in segments}
    by_text = {s.text: layout[s.segment_id] for s in segments}
    if case in ("columns", "mirrored_reordered"):
        assert by_text["막처리 A"]["box_mm"][0] == by_text[expected["a"]]["box_mm"][0]
        assert by_text["세정 B"]["box_mm"][0] == by_text[expected["b"]]["box_mm"][0]
        assert (by_text["막처리 A"]["box_mm"][0] < by_text["세정 B"]["box_mm"][0]) == (case == "columns")
        assert expected["event"] in by_text
    elif case == "ambiguous":
        assert by_text["막처리 A"] == by_text["세정 B"]
        assert expected["ambiguous"] in by_text  # no automatic label assignment
    else:
        assert all(text in by_text for text in expected["conflict"])
    assert [(s.segment_id, s.locator, s.text) for s in segments] == before
    assert all(set(v) == {"slide", "box_mm"} for v in layout.values())
