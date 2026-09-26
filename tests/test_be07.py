"""BE-07 확인: 렌더 어댑터(PDF/DOCX)·스냅샷 고정·findings/not_checked·출력 식별값 공유(㉖)·템플릿 가드.

모든 자료는 가상(fixture mock 묶음·Pillow 생성 PNG). 실제 회사 자료·AI 호출 없음.
PDF는 시스템 Chromium 계열 브라우저가 있을 때만 실제로 만든다. 없으면 해당 테스트는 skip으로 표시된다(통과가 아니다).
"""
from __future__ import annotations

import hashlib
import inspect
import io
import json
import os
import subprocess
import threading
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app import create_app
from app.config import Settings
from app.db import connect
from app.models import Block, Document, Page
from app.services import export_render as er
from app.services import layout_checks
from app.services.documents import get_current

FIX = Path(__file__).parent / "fixtures" / "ddalgi_mock_bundle_v1"
BRIEF = {"purpose": "테스트", "emphasis": [], "direction": "balanced", "target_pages": 4, "photo_preference": "balanced"}
CLEAN_TXT = "회사명: 예시 회사\n회사 개요: 예시 회사는 가상 부품 표면처리와 검사를 하는 테스트 기업입니다.\n사업 분야: 가상 부품 표면처리\n".encode()
LONG_UNIT = "[MOCK] 긴 문단의 줄바꿈과 배치 경고를 시험하는 가상 문장입니다."
BANNED = ("거산", "케미칼", "Geosan")

# 템플릿·폰트·DOCX 배치 상수·PDF 렌더 상수의 sha256(줄바꿈 정규화). 이 중 하나라도 바꾸면 TEMPLATE_VERSION을 올리고 여기 값을 갱신한다.
TEMPLATE_FINGERPRINTS = {
    "template_v0": "7b1b3eaabcad9c23078f68a09fd2ccba89a372cbaec9b13643ebfc3abbbe8096",
}

# 브라우저 탐색은 어댑터와 같은 규칙(EXPORT_BROWSER_PATH 우선). 테스트 settings에도 같은 값을 넣는다.
BROWSER_PATH_ENV = (os.environ.get("EXPORT_BROWSER_PATH") or "").strip() or None
BROWSER = er.find_browser(Settings(private_runs_dir=Path("."), db_path=Path("."), export_browser_path=BROWSER_PATH_ENV))
needs_browser = pytest.mark.skipif(BROWSER is None, reason="Chromium 계열 브라우저 없음 — PDF 렌더 테스트 미실행")


def _png(width: int = 8, height: int = 6, color=(10, 20, 30)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buf, format="PNG")
    return buf.getvalue()


def _jpeg(width: int, height: int, orientation: int | None = None, color=(200, 40, 40)) -> bytes:
    """왼쪽 위 사분면을 다른 색으로 칠한 JPEG. orientation을 주면 EXIF Orientation 태그를 넣는다(픽셀은 그대로)."""
    from PIL import Image

    img = Image.new("RGB", (width, height), color)
    for x in range(width // 2):
        for y in range(height // 2):
            img.putpixel((x, y), (20, 20, 200))
    buf = io.BytesIO()
    if orientation:
        exif = Image.Exif()
        exif[er.EXIF_ORIENTATION_TAG] = orientation
        img.save(buf, format="JPEG", quality=90, exif=exif.tobytes())
    else:
        img.save(buf, format="JPEG", quality=90)
    return buf.getvalue()


def _blk(bid: str, typ: str, **content) -> Block:
    return Block(block_id=bid, type=typ, content=content)


def _doc(pages: list[Page], *, did: str = "doc_t", rev: int = 1, title: str = "[MOCK] 테스트 문서") -> Document:
    return Document(document_id=did, session_id="sess_t", document_revision=rev, input_revision=1, title=title,
                    target_pages=4, status="draft", pages=pages)


def _fixture_doc(name: str) -> Document:
    return Document.model_validate(json.loads((FIX / "ui" / "documents" / f"document_{name}.json").read_text(encoding="utf-8")))


def _fixture_snapshot(name: str) -> er.RenderSnapshot:
    doc = _fixture_doc(name)
    assets = {aid: er.asset_from_bytes(aid, (FIX / "ingest" / "images" / f"{aid}.png").read_bytes())
              for aid in layout_checks.image_asset_ids(doc)}
    return er.snapshot_from_document(doc, assets)


def _stress_snapshot() -> er.RenderSnapshot:
    """넘침·가로/세로 사진·깨진 이미지·없는 asset·사진 자리·HTML 문자를 한 문서에."""
    pages = [
        Page(page_id="p1", title="표지", layout_key="text_photo", blocks=[
            _blk("h1", "heading", text="[MOCK] 스트레스 문서 한글 서체 확인", level=1),
            _blk("para", "paragraph", text="<script>alert(1)</script> & <img src=x onerror=alert(2)> \"따옴표\""),
            _blk("lst", "list", items=["가상 공정 A", "가상 공정 B"]),
            _blk("img_land", "image", asset_id="LAND", alt="가로", caption="[MOCK] 가로", fit="contain"),
        ]),
        Page(page_id="p2", title="긴 본문", layout_key="text", blocks=[
            _blk("h2", "heading", text="긴 본문", level=2),
            _blk("long", "paragraph", text=" ".join([LONG_UNIT] * 150)),
        ]),
        Page(page_id="p3", title="사진", layout_key="text_photo", blocks=[
            _blk("img_port", "image", asset_id="PORT", alt="세로", caption="[MOCK] 세로", fit="contain"),
            _blk("img_crop", "image", asset_id="WIDE", alt="배너", caption="[MOCK] 배너", fit="crop"),
            _blk("img_broken", "image", asset_id="BROKEN", alt="", caption="[MOCK] 깨진 파일", fit="contain"),
            _blk("img_missing", "image", asset_id="NOPE", alt="", caption="", fit="contain"),
            _blk("ph", "image_placeholder", description="[MOCK] 사진을 추가할 자리"),
        ]),
    ]
    assets = {"LAND": er.asset_from_bytes("LAND", _png(96, 64)), "PORT": er.asset_from_bytes("PORT", _png(64, 96)),
              "WIDE": er.asset_from_bytes("WIDE", _png(128, 42)), "BROKEN": er.asset_from_bytes("BROKEN", b"\x89PNG\r\n\x1a\nbroken" * 5)}
    return er.snapshot_from_document(_doc(pages, did="doc_stress"), assets)


def _ok_assets(snap: er.RenderSnapshot) -> int:
    return sum(1 for a in snap.assets.values() if a.ok)


@pytest.fixture
def settings(tmp_path):
    return Settings(private_runs_dir=tmp_path / "runs", db_path=tmp_path / "runs" / "t.sqlite3", export_browser_path=BROWSER_PATH_ENV)


@pytest.fixture
def app(settings):
    return create_app(settings)


@pytest.fixture
def out_dir(tmp_path):
    return tmp_path / "runs" / "exports"


class Flow:
    """세션 + 가상 TXT/PNG 업로드 + mock 초안 rev.1(이미지 블록 포함)."""

    def __init__(self, app, *, png: bytes | None = None):
        self.c = TestClient(app)
        self.sid = self.c.post("/api/v1/sessions", json={"brief": BRIEF}).json()["session_id"]
        files = [("files", ("a.txt", io.BytesIO(CLEAN_TXT))), ("files", ("p.png", io.BytesIO(png or _png())))]
        up = self.c.post(f"/api/v1/sessions/{self.sid}/sources", files=files).json()
        selected = [i["source_id"] for i in up["items"]]
        r = self.c.patch(f"/api/v1/sessions/{self.sid}/inputs", json={"expected_input_revision": 1, "selected_source_ids": selected})
        assert r.status_code == 200, r.text
        self.rev_in = r.json()["input_revision"]
        job = self.c.post(f"/api/v1/sessions/{self.sid}/preflights", json={"expected_input_revision": self.rev_in}).json()
        pf = self._job(job["job_id"])["result_ref"]["preflight_id"]
        job = self.c.post(f"/api/v1/sessions/{self.sid}/drafts", json={"preflight_id": pf, "input_revision": self.rev_in, "confirmed": True}).json()
        self.did = self._job(job["job_id"])["result_ref"]["document_id"]

    def _job(self, jid):
        job = self.c.get(f"/api/v1/sessions/{self.sid}/jobs/{jid}").json()
        assert job["status"] == "succeeded", job
        return job


# ================= 출력 식별값 공유(㉖) =================

def test_identity_values_come_from_layout_checks(out_dir):
    r = er.render(_fixture_snapshot("1pages"), "docx", out_dir)
    assert r.template_version == layout_checks.TEMPLATE_VERSION == "template_v0"
    assert r.render_options_hash == layout_checks.RENDER_OPTIONS_HASH == layout_checks.render_options_hash(layout_checks.DEFAULT_RENDER_OPTIONS)
    assert len(r.render_options_hash) == 16 and int(r.render_options_hash, 16) >= 0
    changed = dict(layout_checks.DEFAULT_RENDER_OPTIONS, margin_mm=20)
    assert layout_checks.render_options_hash(changed) != r.render_options_hash
    # 옵션 내용이 실제 값(용지·여백·폰트)인지 — 'default' 자리표시자가 아니다
    assert layout_checks.DEFAULT_RENDER_OPTIONS["font_family"] == "Pretendard" and layout_checks.DEFAULT_RENDER_OPTIONS["page_size"] == "A4"


def test_adapter_reads_layout_checks_at_call_time(out_dir, monkeypatch):
    """어댑터가 값을 복사해 두지 않고 layout_checks의 현재 값을 쓴다(한쪽만 올리면 바로 드러남)."""
    monkeypatch.setattr(layout_checks, "TEMPLATE_VERSION", "template_vX")
    monkeypatch.setattr(layout_checks, "RENDER_OPTIONS_HASH", "0123456789abcdef")
    r = er.render(_fixture_snapshot("1pages"), "docx", out_dir)
    assert r.template_version == "template_vX" and r.render_options_hash == "0123456789abcdef"


def test_manifest_hash_pure_function_matches_legacy_formula():
    items = [("b", "h2"), ("a", "h1"), ("a", "")]
    legacy = hashlib.sha256(json.dumps(sorted(items)).encode()).hexdigest()[:16]
    assert layout_checks.manifest_hash(items) == legacy
    assert layout_checks.manifest_hash([]) == hashlib.sha256(b"[]").hexdigest()[:16]


def test_template_fingerprint_pinned_per_version():
    """템플릿·폰트·DOCX 배치 상수·PDF 렌더 상수를 바꾸면서 TEMPLATE_VERSION을 올리지 않으면 실패한다."""
    pinned = TEMPLATE_FINGERPRINTS.get(layout_checks.TEMPLATE_VERSION)
    assert pinned, f"TEMPLATE_VERSION {layout_checks.TEMPLATE_VERSION}의 지문이 TEMPLATE_FINGERPRINTS에 없습니다"
    actual = er.template_fingerprint()
    assert actual == pinned, ("템플릿/폰트/DOCX 배치 상수/PDF 렌더 상수가 바뀌었습니다. layout_checks.TEMPLATE_VERSION을 올리고 "
                              f"TEMPLATE_FINGERPRINTS[새 버전] = '{actual}' 을 추가하세요.")
    # 줄바꿈(autocrlf)에 흔들리지 않는다
    raw = er.TEMPLATE_FILE.read_bytes()
    assert b"\r\n" not in raw
    src = inspect.getsource(er.template_fingerprint)
    assert "PDF_RENDER_CONSTANTS" in src and "DOCX_LAYOUT_CONSTANTS" in src


def test_render_takes_no_caller_options():
    params = inspect.signature(er.render).parameters
    assert list(params) == ["snapshot", "fmt", "out_dir", "settings"]
    assert all(p.kind == p.POSITIONAL_OR_KEYWORD for p in params.values())


# ================= 스냅샷 고정 =================

def test_snapshot_manifest_matches_db_and_tracks_content_hash(app, settings, out_dir):
    flow = Flow(app)
    with connect(settings.db_path) as conn:
        doc = get_current(conn, flow.sid, flow.did)
        assert layout_checks.image_asset_ids(doc), "mock 초안에 image 블록이 있어야 한다"
        snap = er.build_snapshot(conn, settings, flow.sid, doc)
        assert snap.asset_manifest_hash == layout_checks.asset_manifest_hash(conn, doc)
        aid = layout_checks.image_asset_ids(doc)[0]
        asset = snap.assets[aid]
        assert asset.ok and asset.data and asset.width == 8 and asset.height == 6 and len(asset.content_hash) == 64
        # 같은 이름으로 파일을 바꿔치기한 상황: DB content_hash와 파일이 어긋남 → 스냅샷은 hash_mismatch, manifest는 DB값 기준으로 둘 다 동일하게 변함
        conn.execute("UPDATE assets SET content_hash='changed' WHERE asset_id=?", (aid,))
        snap2 = er.build_snapshot(conn, settings, flow.sid, doc)
        assert snap2.assets[aid].ok is False and snap2.assets[aid].reason == "hash_mismatch"
        assert snap2.asset_manifest_hash == layout_checks.asset_manifest_hash(conn, doc) != snap.asset_manifest_hash
    r = er.render(snap2, "docx", out_dir)
    assert r.asset_manifest_hash == snap2.asset_manifest_hash
    assert [f.kind for f in r.findings] == ["broken_image"] and r.findings[0].details["reason"] == "hash_mismatch"


def test_snapshot_access_rules_missing_file_and_not_ready(app, settings):
    flow_a, flow_b = Flow(app), Flow(app)
    with connect(settings.db_path) as conn:
        doc_a = get_current(conn, flow_a.sid, flow_a.did)
        doc_b = get_current(conn, flow_b.sid, flow_b.did)
        aid_b = layout_checks.image_asset_ids(doc_b)[0]
        # A 문서가 B 세션의 asset을 가리키면 not_accessible(존재를 쓰지 않는다). manifest는 DB content_hash로 계산돼 승인 검사와 같다
        pages = [Page(page_id="p", title="t", layout_key="text_photo", blocks=[
            _blk("x", "image", asset_id=aid_b, alt="", caption="", fit="contain")])]
        foreign = doc_a.model_copy(update={"pages": pages})
        snap = er.build_snapshot(conn, settings, flow_a.sid, foreign)
        assert snap.assets[aid_b].ok is False and snap.assets[aid_b].reason == "not_accessible" and snap.assets[aid_b].data is None
        assert snap.asset_manifest_hash == layout_checks.asset_manifest_hash(conn, foreign)
        # 아직 준비되지 않은 asset → not_ready, 바이트 없음
        aid_a = layout_checks.image_asset_ids(doc_a)[0]
        conn.execute("UPDATE assets SET status='processing' WHERE asset_id=?", (aid_a,))
        snap_nr = er.build_snapshot(conn, settings, flow_a.sid, doc_a)
        assert snap_nr.assets[aid_a].reason == "not_ready" and snap_nr.assets[aid_a].data is None
        conn.execute("UPDATE assets SET status='ready' WHERE asset_id=?", (aid_a,))
        # 파일이 사라진 asset → missing
        stored = conn.execute("SELECT stored_path FROM assets WHERE asset_id=?", (aid_a,)).fetchone()[0]
        (settings.private_runs_dir / stored).unlink()
        snap_a = er.build_snapshot(conn, settings, flow_a.sid, doc_a)
        assert snap_a.assets[aid_a].reason == "missing"
        # 아예 없는 asset_id → not_found, content_hash ""
        ghost = doc_a.model_copy(update={"pages": [Page(page_id="p", title="t", layout_key="text", blocks=[
            _blk("g", "image", asset_id="asset_ghost", alt="", caption="", fit="contain")])]})
        snap_g = er.build_snapshot(conn, settings, flow_a.sid, ghost)
        assert snap_g.assets["asset_ghost"].reason == "not_found" and snap_g.manifest_items == [("asset_ghost", "")]


def test_render_uses_snapshot_bytes_not_disk(app, settings, out_dir):
    """규칙 ③: 스냅샷을 만든 뒤 파일·DB 행이 사라져도 렌더는 스냅샷 바이트로 이미지를 넣는다."""
    flow = Flow(app)
    with connect(settings.db_path) as conn:
        doc = get_current(conn, flow.sid, flow.did)
        snap = er.build_snapshot(conn, settings, flow.sid, doc)
        aid = layout_checks.image_asset_ids(doc)[0]
        stored = conn.execute("SELECT stored_path FROM assets WHERE asset_id=?", (aid,)).fetchone()[0]
        conn.execute("DELETE FROM assets WHERE asset_id=?", (aid,))
    (settings.private_runs_dir / stored).unlink()
    r = er.render(snap, "docx", out_dir)
    assert r.findings == [] and er.docx_info(r.file_path)["inline_shapes"] == 1
    if BROWSER is not None:
        rp = er.render(snap, "pdf", out_dir, settings)
        from pypdf import PdfReader

        assert rp.findings == [] and sum(len(p.images) for p in PdfReader(str(rp.file_path)).pages) >= 1


def test_asset_from_bytes_verifies_hash_decodes_and_restricts_format():
    data = _png(5, 7)
    ok = er.asset_from_bytes("a", data)
    assert ok.ok and (ok.width, ok.height) == (5, 7) and ok.content_hash == hashlib.sha256(data).hexdigest() and ok.mime_type == "image/png"
    assert er.asset_from_bytes("a", data, content_hash="0" * 64).reason == "hash_mismatch"
    assert er.asset_from_bytes("a", b"not an image").reason == "decode_failed"
    assert er.asset_from_bytes("a", None).reason == "missing"
    # Pillow는 열지만 두 렌더러가 모두 지원한다고 볼 수 없는 형식(BMP/WebP 등) → unsupported_format, 바이트 없음
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (3, 3)).save(buf, format="BMP")
    bmp = er.asset_from_bytes("b", buf.getvalue())
    assert bmp.ok is False and bmp.reason == "unsupported_format" and bmp.data is None
    snap = er.snapshot_from_document(_doc([Page(page_id="p", title="t", layout_key="text", blocks=[
        _blk("i", "image", asset_id="unknown", alt="", caption="", fit="contain")])]), {})
    assert snap.assets["unknown"].reason == "not_found"


def test_embed_bytes_downscales_only_large_images():
    small = er.asset_from_bytes("s", _png(300, 200))
    data, w, h = er.embed_bytes(small)
    assert data is small.data and (w, h) == (300, 200)          # 상한 이하: 검증한 바이트 그대로
    big = er.asset_from_bytes("b", _png(er.IMAGE_MAX_PX * 2, 400))
    data, w, h = er.embed_bytes(big)
    assert w == er.IMAGE_MAX_PX and h == 200 and data != big.data
    assert big.content_hash == hashlib.sha256(big.data).hexdigest()   # 스냅샷 자체는 바뀌지 않는다
    from PIL import Image

    with Image.open(io.BytesIO(data)) as img:
        assert img.format == "PNG" and img.size == (er.IMAGE_MAX_PX, 200)


# ================= DOCX =================

def test_docx_render_structure_checks_and_font(out_dir):
    snap = _fixture_snapshot("10pages")
    r = er.render(snap, "docx", out_dir)
    assert r.format == "docx" and r.file_path.is_file() and r.file_path.stat().st_size > 0
    assert r.actual_pages is None                       # 2.3절: 논리 페이지 수·PDF 쪽수를 대신 넣지 않는다
    assert r.layout_ok is False and r.not_checked == ["overflow"]
    overflow = next(c for c in r.checks if c.check_key == "overflow")
    assert overflow.required is True and overflow.result == "not_checked" and overflow.reason == "docx_no_layout_engine"
    assert {c.check_key: c.result for c in r.checks if c.check_key != "overflow"} == {"broken_image": "ok", "placeholder_remaining": "ok"}
    assert r.findings == [] and r.renderer.startswith("python-docx/")
    # 구조만 다시 센다(쪽수 판단에 쓰지 않는다)
    info = er.docx_info(r.file_path)
    total_blocks = sum(len(p.blocks) for p in snap.pages)
    assert info["inline_shapes"] == _ok_assets(snap) == 2 and info["paragraphs"] >= total_blocks and info["sections"] == 1
    assert r.details["font_embedding"] == "not_supported_in_this_implementation"
    assert sorted(p.name for p in out_dir.iterdir()) == [r.file_path.name]   # 임시 파일이 남지 않는다
    # 글꼴 이름: Normal·Heading·Caption 스타일에 eastAsia 포함 지정, 테마 글꼴 속성 제거
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(str(r.file_path))
    for name in ("Normal", "Heading 1", "Heading 2", "Caption", "List Bullet"):
        rfonts = d.styles[name].element.rPr.find(qn("w:rFonts"))
        assert rfonts.get(qn("w:eastAsia")) == "Pretendard" == rfonts.get(qn("w:ascii")), name
        assert rfonts.get(qn("w:eastAsiaTheme")) is None and rfonts.get(qn("w:asciiTheme")) is None, name
    assert d.core_properties.title == snap.title and d.core_properties.author == ""
    sec = d.sections[0]
    assert round(sec.page_width.mm) == 210 and round(sec.page_height.mm) == 297 and round(sec.left_margin.mm) == 15
    assert "[MOCK] 예시 회사 소개" in "\n".join(p.text for p in d.paragraphs)


def test_docx_findings_broken_missing_placeholder_and_escaping(out_dir):
    snap = _stress_snapshot()
    r = er.render(snap, "docx", out_dir)
    kinds = sorted((f.kind, f.block_id) for f in r.findings)
    assert kinds == [("broken_image", "img_broken"), ("broken_image", "img_missing"), ("placeholder_remaining", "ph")]
    by_key = {c.check_key: c for c in r.checks}
    assert by_key["broken_image"].result == "finding" and by_key["broken_image"].block_ids == ["img_broken", "img_missing"]
    assert by_key["placeholder_remaining"].result == "finding" and by_key["placeholder_remaining"].page_ids == ["p3"]
    assert r.layout_ok is False
    import docx

    d = docx.Document(str(r.file_path))
    text = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "[사진 자리]" in text and "[이미지를 열 수 없음]" in text and "decode_failed" in text and "not_found" in text
    assert "<script>alert(1)</script>" in text          # 문자 그대로. 실행·해석되지 않는다
    assert er.docx_info(r.file_path)["inline_shapes"] == _ok_assets(snap) == 3   # 정상 이미지만 삽입(가로·세로·배너)
    assert not any(w in text for w in BANNED)


def test_docx_xml_incompatible_chars_are_stripped_not_crashing(out_dir):
    """U+FFFE/U+FFFF·서로게이트는 lxml이 거부한다. 제거해서 만들고, 그래도 실패하면 RenderError로만 알린다."""
    pages = [Page(page_id="p", title="t\ufffe", layout_key="text", blocks=[
        _blk("h", "heading", text="제목\uffff", level=1), _blk("q", "paragraph", text="a\ufffeb\x00c")])]
    snap = er.snapshot_from_document(_doc(pages, title="제목\ufffe"), {})
    r = er.render(snap, "docx", out_dir)
    import docx

    d = docx.Document(str(r.file_path))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "abc" in text and "제목" in text and "\ufffe" not in text and d.core_properties.title == "제목"


def test_docx_build_failure_is_render_error_without_partial_file(out_dir, monkeypatch):
    def boom(*a, **k):
        raise ValueError("boom")

    monkeypatch.setattr(er, "_build_docx", boom)
    with pytest.raises(er.RenderError) as e:
        er.render(_fixture_snapshot("1pages"), "docx", out_dir)
    assert e.value.code == "render_failed" and not any(out_dir.iterdir())
    monkeypatch.undo()
    import docx.document

    def disk_full(self, path):
        raise OSError("disk full")

    monkeypatch.setattr(docx.document.Document, "save", disk_full)
    with pytest.raises(er.RenderError) as e:
        er.render(_fixture_snapshot("1pages"), "docx", out_dir)
    assert e.value.code == "save_failed" and not any(out_dir.iterdir())


def test_docx_concurrent_renders_of_same_revision_do_not_collide(out_dir):
    snap = _fixture_snapshot("1pages")
    results, errors = [], []

    def run():
        try:
            results.append(er.render(snap, "docx", out_dir))
        except Exception as exc:  # noqa: BLE001
            errors.append(exc)
    threads = [threading.Thread(target=run) for _ in range(3)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert errors == [] and len(results) == 3
    assert sorted(p.name for p in out_dir.iterdir()) == [results[0].file_path.name]
    assert er.docx_info(results[0].file_path)["inline_shapes"] == 1


# ================= HTML(템플릿) 안전성 — 브라우저 없이 확인 =================

def test_html_escapes_untrusted_text_and_has_no_external_resources():
    html = er.build_html(_stress_snapshot())
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "<img src=x" not in html and "&lt;img src=x onerror=alert(2)&gt;" in html   # 태그가 아니라 문자로만 남는다
    assert html.count("<script") == 1 and 'nonce="' in html          # 측정 스크립트 하나뿐
    assert "http://" not in html and "https://" not in html and "file://" not in html   # 외부 리소스 없음(폰트·이미지는 data URI)
    assert "default-src 'none'" in html and 'font-family: "Pretendard"' in html
    assert "[사진 자리]" in html and "[이미지를 열 수 없음]" in html
    assert 'data-block-id="long"' in html and 'data-page-id="p2"' in html
    assert "document.fonts.ready" in html   # 측정은 폰트 로드 뒤


def test_measure_json_validation_rejects_forged_or_foreign_ids():
    snap = _fixture_snapshot("1pages")
    good = {"fonts_ready": True, "pages": [{"page_id": "page_mock_single", "height_px": 10.0, "overflow": False, "excess_px": 0, "first_overflow_block_id": None}]}
    assert er._valid_measure(good, snap) is True
    bad_cases = [
        dict(good, fonts_ready=False),                                                        # 폰트 미로드
        {"fonts_ready": True, "pages": []},                                                   # 페이지 누락
        {"fonts_ready": True, "pages": [dict(good["pages"][0], page_id="other")]},            # 스냅샷에 없는 page_id
        {"fonts_ready": True, "pages": [dict(good["pages"][0], overflow="false")]},           # 타입 위조
        {"fonts_ready": True, "pages": [dict(good["pages"][0], first_overflow_block_id="nope")]},  # 없는 block_id
        {"fonts_ready": True, "pages": good["pages"] + good["pages"]},                        # 중복
        "[]", None,
    ]
    for bad in bad_cases:
        assert er._valid_measure(bad, snap) is False, bad
    # print_pdf_and_measure는 unescape하지 않는다(ID 안의 &…; 로 JSON 구조를 위조하지 못하게)
    assert "unescape(" not in inspect.getsource(er.print_pdf_and_measure)


# ================= PDF (브라우저 필요) =================

@needs_browser
def test_pdf_render_pages_fonts_text(settings, out_dir):
    r = er.render(_fixture_snapshot("1pages"), "pdf", out_dir, settings)
    assert r.file_path.suffix == ".pdf" and r.file_path.stat().st_size > 0
    assert r.actual_pages == 1 and r.layout_ok is True and r.not_checked == [] and r.findings == []
    assert all(c.required and c.result == "ok" for c in r.checks)
    assert r.renderer.split("/")[0] in ("chrome", "edge", "chromium") and r.renderer.split("/")[1] not in ("", "unknown")
    assert any("Pretendard" in f for f in r.details["fonts"]), r.details["fonts"]
    assert r.details["measure"]["fonts_ready"] is True
    info = er.pdf_info(r.file_path)
    assert "예시 회사" in info["first_page_text"] and "가상 공정 A" in info["first_page_text"]
    assert not any(w in info["first_page_text"] for w in BANNED)
    # 임시 HTML·프로필은 남지 않는다
    assert sorted(p.name for p in out_dir.iterdir()) == [r.file_path.name]
    r10 = er.render(_fixture_snapshot("10pages"), "pdf", out_dir, settings)
    assert r10.actual_pages == 10 and r10.layout_ok is True


@needs_browser
def test_pdf_overflow_broken_placeholder_images(settings, out_dir):
    snap = _stress_snapshot()
    r = er.render(snap, "pdf", out_dir, settings)
    by_kind = {}
    for f in r.findings:
        by_kind.setdefault(f.kind, []).append(f)
    assert [f.page_id for f in by_kind["overflow"]] == ["p2"] and by_kind["overflow"][0].block_id == "long"
    assert by_kind["overflow"][0].details["excess_mm"] > 0
    assert sorted(f.block_id for f in by_kind["broken_image"]) == ["img_broken", "img_missing"]
    assert [f.block_id for f in by_kind["placeholder_remaining"]] == ["ph"]
    assert r.layout_ok is False and r.not_checked == []
    assert r.actual_pages > len(snap.pages)              # 넘친 페이지는 잘리지 않고 다음 물리 쪽으로 흐른다
    checks = {c.check_key: c for c in r.checks}
    assert checks["overflow"].result == "finding" and checks["overflow"].page_ids == ["p2"]
    from pypdf import PdfReader

    reader = PdfReader(str(r.file_path))
    text = "\n".join(p.extract_text() or "" for p in reader.pages)
    assert "<script>alert(1)</script>" in text and "[사진 자리]" in text and "[이미지를 열 수 없음]" in text
    images = sum(len(p.images) for p in reader.pages)
    assert images >= _ok_assets(snap) == 3                # 정상 이미지(가로·세로·배너)는 모두 실제 이미지로 삽입
    measured = {p["page_id"]: p for p in r.details["measure"]["pages"]}
    assert measured["p1"]["overflow"] is False and measured["p2"]["overflow"] is True


@needs_browser
def test_pdf_measure_failed_is_not_checked_but_file_and_pages_are_real(settings, out_dir, monkeypatch):
    original = er.print_pdf_and_measure

    def no_measure(*a, **k):
        original(*a, **k)
        return None
    monkeypatch.setattr(er, "print_pdf_and_measure", no_measure)
    r = er.render(_fixture_snapshot("1pages"), "pdf", out_dir, settings)
    assert r.actual_pages == 1 and r.file_path.is_file()
    assert r.not_checked == ["overflow"] and r.layout_ok is False
    assert next(c for c in r.checks if c.check_key == "overflow").reason == "measure_failed"

    def bad_measure(*a, **k):
        original(*a, **k)
        return {"fonts_ready": True, "pages": []}   # 위조·불일치 → measure_invalid
    monkeypatch.setattr(er, "print_pdf_and_measure", bad_measure)
    r2 = er.render(_fixture_snapshot("1pages"), "pdf", out_dir, settings)
    assert r2.layout_ok is False and next(c for c in r2.checks if c.check_key == "overflow").reason == "measure_invalid"


@needs_browser
def test_pdf_render_errors_leave_no_partial_file(settings, out_dir, monkeypatch):
    def timeout(cmd, timeout, what):
        raise er.RenderError("render_timeout", "t")

    monkeypatch.setattr(er, "_run", timeout)
    with pytest.raises(er.RenderError) as e:
        er.render(_fixture_snapshot("1pages"), "pdf", out_dir, settings)
    assert e.value.code == "render_timeout"
    assert not out_dir.exists() or list(out_dir.iterdir()) == []
    # 브라우저가 0이 아닌 코드로 끝나고 PDF가 없으면 render_failed
    monkeypatch.setattr(er, "_run", lambda cmd, timeout, what: subprocess.CompletedProcess(cmd, 1, b"", b"crash"))
    with pytest.raises(er.RenderError) as e:
        er.render(_fixture_snapshot("1pages"), "pdf", out_dir, settings)
    assert e.value.code == "render_failed" and list(out_dir.iterdir()) == []


def test_render_errors_without_browser(tmp_path, out_dir):
    with pytest.raises(er.RenderError) as e:
        er.render(_fixture_snapshot("1pages"), "pptx", out_dir)
    assert e.value.code == "unsupported_format"
    no_browser = Settings(private_runs_dir=tmp_path / "runs", db_path=tmp_path / "runs" / "t.sqlite3",
                          export_browser_path=str(tmp_path / "nope" / "chrome.exe"))
    assert er.find_browser(no_browser) is None
    with pytest.raises(er.RenderError) as e:
        er.render(_fixture_snapshot("1pages"), "pdf", out_dir, no_browser)
    assert e.value.code == "browser_not_found"
    assert not out_dir.exists() or list(out_dir.iterdir()) == []


# ================= 규칙 =================

def test_required_not_checked_is_never_layout_ok():
    findings: list = []
    checks = [er._record("overflow", "docx", findings, not_checked_reason="docx_no_layout_engine"),
              er._record("broken_image", "docx", findings), er._record("placeholder_remaining", "docx", findings)]
    assert er._layout_ok(checks) is False and all(c.required for c in checks)
    assert er._layout_ok([er._record(k, "pdf", findings) for k in er.CHECK_KEYS]) is True
    assert er.REQUIRED_CHECKS["pdf"] == er.REQUIRED_CHECKS["docx"] == frozenset(er.CHECK_KEYS)


def test_browser_args_have_no_shell_and_layout_flags_are_fingerprinted(tmp_path):
    args = er._browser_args(Path("C:/x/chrome.exe"), tmp_path / "prof")
    assert Path(args[0]) == Path("C:/x/chrome.exe") and "--headless=new" in args and any(a.startswith("--user-data-dir=") for a in args)
    assert f"--window-size={er.PDF_RENDER_CONSTANTS['window_size']}" in args
    assert f"--virtual-time-budget={er.PDF_RENDER_CONSTANTS['virtual_time_budget_ms']}" in args
    assert "shell=True" not in inspect.getsource(er._run)


def test_no_real_company_terms_in_export_code_and_template():
    sources = [inspect.getsource(er), inspect.getsource(layout_checks), er.TEMPLATE_FILE.read_text(encoding="utf-8"),
               (er.FONT_DIR / "SOURCE.md").read_text(encoding="utf-8")]
    for src in sources:
        for banned in BANNED:
            assert banned not in src


# ================= 회귀: 스냅샷 격리·전체 디코딩·EXIF 방향·샌드박스 =================

def test_snapshot_is_isolated_from_document_mutation(app, settings, out_dir):
    """규칙: 스냅샷을 만든 뒤 원본 Document(페이지·블록·content dict)를 바꿔도 스냅샷·해시·출력은 그대로다."""
    flow = Flow(app)
    with connect(settings.db_path) as conn:
        doc = get_current(conn, flow.sid, flow.did)
        snap_db = er.build_snapshot(conn, settings, flow.sid, doc)
    doc_local = _fixture_doc("1pages")
    snap_local = er.snapshot_from_document(doc_local, {aid: er.asset_from_bytes(aid, (FIX / "ingest" / "images" / f"{aid}.png").read_bytes())
                                                       for aid in layout_checks.image_asset_ids(doc_local)})
    for snap, original in ((snap_db, doc), (snap_local, doc_local)):
        before_pages = [pg.model_dump() for pg in snap.pages]
        before_hash, before_manifest = snap.content_hash, snap.asset_manifest_hash
        before_html = er.build_html(snap)
        # 원본을 여러 방식으로 변경: content dict 직접 수정, 블록 추가, 페이지 제목 변경, 페이지 삭제, image asset_id 교체
        original.pages[0].blocks[0].content["text"] = "변경된 제목"
        original.pages[0].blocks.append(_blk("added", "paragraph", text="추가 블록"))
        original.pages[0].title = "바뀐 페이지"
        for b in original.pages[0].blocks:
            if b.type == "image":
                b.content["asset_id"] = "asset_swapped"
        del original.pages[-1:]
        assert [pg.model_dump() for pg in snap.pages] == before_pages
        assert snap.content_hash == before_hash and snap.asset_manifest_hash == before_manifest
        after_html = er.build_html(snap)
        assert "변경된 제목" not in after_html and "추가 블록" not in after_html and "asset_swapped" not in after_html
        assert len(after_html) == len(before_html)
    r = er.render(snap_local, "docx", out_dir)
    import docx

    text = "\n".join(pg.text for pg in docx.Document(str(r.file_path)).paragraphs)
    assert "[MOCK] 예시 회사 소개" in text and "변경된 제목" not in text and "추가 블록" not in text
    assert r.findings == [] and er.docx_info(r.file_path)["inline_shapes"] == 1


def test_truncated_jpeg_is_decode_failed_and_broken_image(out_dir):
    """헤더는 정상이지만 본문이 잘린 JPEG: 전체 픽셀 디코딩에서 걸려 decode_failed → 출력에는 깨진 이미지 상자."""
    full = _jpeg(240, 160)
    truncated = full[: int(len(full) * 0.6)]
    ok = er.asset_from_bytes("full", full)
    assert ok.ok and (ok.width, ok.height) == (240, 160) and ok.mime_type == "image/jpeg"
    bad = er.asset_from_bytes("cut", truncated)
    assert bad.ok is False and bad.reason == "decode_failed" and bad.data is None
    # 잘린 파일도 Pillow.open은 성공한다(헤더만) — 검증이 헤더에 그치면 놓치는 사례임을 같이 확인
    from PIL import Image

    with Image.open(io.BytesIO(truncated)) as img:
        assert img.size == (240, 160)
    pages = [Page(page_id="p", title="t", layout_key="text_photo", blocks=[
        _blk("i_cut", "image", asset_id="cut", alt="", caption="[MOCK] 잘린 JPEG", fit="contain"),
        _blk("i_ok", "image", asset_id="full", alt="", caption="[MOCK] 정상 JPEG", fit="contain")])]
    snap = er.snapshot_from_document(_doc(pages), {"cut": bad, "full": ok})
    r = er.render(snap, "docx", out_dir)
    assert [(f.kind, f.block_id, f.details["reason"]) for f in r.findings] == [("broken_image", "i_cut", "decode_failed")]
    assert er.docx_info(r.file_path)["inline_shapes"] == 1 and r.layout_ok is False
    html = er.build_html(snap)
    assert html.count("data:image/jpeg;base64,") == 1 and "[이미지를 열 수 없음]" in html and "decode_failed" in html


@pytest.mark.parametrize("size", [(400, 300), (3200, 2400)], ids=["under_limit", "over_limit"])
def test_embed_bytes_applies_exif_orientation_and_keeps_snapshot_bytes(size):
    """EXIF Orientation=6(시계 방향 90°)인 가로 JPEG는 세로로 삽입된다. 상한 이하·초과 모두 비율 유지, 스냅샷 바이트·해시 보존."""
    from PIL import Image

    width, height = size
    raw = _jpeg(width, height, orientation=6)
    asset = er.asset_from_bytes("exif", raw)
    assert asset.ok and (asset.width, asset.height) == (width, height)       # 스냅샷 크기는 파일의 픽셀 크기 그대로(DB와 같은 기준)
    data, w, h = er.embed_bytes(asset)
    expect_w, expect_h = (height, width) if max(size) <= er.IMAGE_MAX_PX else (int(height * er.IMAGE_MAX_PX / max(size)), er.IMAGE_MAX_PX)
    assert (w, h) == (expect_w, expect_h) and abs(w / h - height / width) < 0.01   # 세로 방향, 비율 유지
    assert data != raw and asset.data == raw and asset.content_hash == hashlib.sha256(raw).hexdigest()   # 원본 바이트·해시 보존
    with Image.open(io.BytesIO(data)) as img:
        assert img.format == "JPEG" and img.size == (w, h)
        assert img.getexif().get(er.EXIF_ORIENTATION_TAG, 1) == 1                # 돌린 사본에는 방향 태그가 남지 않는다(이중 회전 방지)
        # 원본의 왼쪽 위(파란 사분면)는 90° 회전 뒤 오른쪽 위로 간다
        assert img.getpixel((w - 2, 1))[2] > 150 and img.getpixel((1, h - 2))[0] > 150
    # 같은 입력을 두 번 처리하면 출력 바이트가 같다
    assert hashlib.sha256(er.embed_bytes(asset)[0]).hexdigest() == hashlib.sha256(data).hexdigest()


def test_embed_bytes_without_exif_keeps_bytes_and_is_deterministic():
    small = er.asset_from_bytes("s", _jpeg(400, 300))
    data, w, h = er.embed_bytes(small)
    assert data is small.data and (w, h) == (400, 300)
    big_png = er.asset_from_bytes("b", _png(er.IMAGE_MAX_PX * 2, 400))
    d1, w1, h1 = er.embed_bytes(big_png)
    d2, w2, h2 = er.embed_bytes(big_png)
    assert (w1, h1) == (w2, h2) == (er.IMAGE_MAX_PX, 200) and hashlib.sha256(d1).hexdigest() == hashlib.sha256(d2).hexdigest()
    assert big_png.data != d1 and big_png.content_hash == hashlib.sha256(big_png.data).hexdigest()
    big_jpg = er.asset_from_bytes("j", _jpeg(er.IMAGE_MAX_PX * 2, 800))
    j1, _, _ = er.embed_bytes(big_jpg)
    j2, _, _ = er.embed_bytes(big_jpg)
    assert hashlib.sha256(j1).hexdigest() == hashlib.sha256(j2).hexdigest() and big_jpg.data != j1


def test_browser_args_never_disable_sandbox(tmp_path, monkeypatch):
    args = er._browser_args(Path("/usr/bin/chromium"), tmp_path / "prof")
    assert not any(a.startswith("--no-sandbox") or a.startswith("--disable-setuid-sandbox") for a in args)
    if hasattr(os, "geteuid"):
        monkeypatch.setattr(os, "geteuid", lambda: 0)
        args_root = er._browser_args(Path("/usr/bin/chromium"), tmp_path / "prof")
        assert "--no-sandbox" not in args_root
    assert '"--no-sandbox"' not in inspect.getsource(er._browser_args)   # 문자열 인자로 존재하지 않는다(주석 언급만 허용)
