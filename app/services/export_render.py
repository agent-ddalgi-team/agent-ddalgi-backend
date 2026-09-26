"""렌더 어댑터(BE-07, D-03) — 승인 스냅샷을 실제 PDF/DOCX 파일로 만든다. BE-08(LayoutCheck Job·Export·다운로드)이 그대로 쓴다.

경계
- 입력은 RenderSnapshot: Document 블록 + 서버가 검증한 이미지 바이트·content_hash·버전을 고정한 값. render()는 DB·파일을 다시 찾지 않는다.
  DB를 만지는 함수는 build_snapshot() 하나뿐이다.
- 출력 식별값(template_version·render_options_hash·asset_manifest_hash)은 app/services/layout_checks.py의 정의를 그대로 쓴다(계약 확인 ㉖).
  render()는 호출자 옵션을 받지 않는다 — 해시에 반영되지 않는 가변 옵션은 없다.
- 검사 3종(overflow / broken_image / placeholder_remaining)은 PDF·DOCX 모두 required. 검사하지 못한 항목은 result=not_checked로 남기고
  layout_ok=False다(빈 findings를 통과로 해석하지 않는다). DOCX는 배치 엔진이 없어 overflow가 not_checked이고 actual_pages=None이다.
- PDF는 시스템 Chromium 계열 브라우저(Chrome/Edge)를 1회 실행해 headless 인쇄(--print-to-pdf)와 DOM 측정(--dump-dom)을 함께 얻는다.
  브라우저 종류·버전은 해시 밖이며 RenderResult.renderer에 기록한다. 검사와 출력 사이에 렌더러가 바뀌었을 때의 정책은 BE-08 몫.
- 문서 값은 Jinja2 autoescape로만 HTML에 들어가며 첨부·문서 안 문구를 HTML/스크립트로 실행하지 않는다. 외부 리소스는 없다.
- 삽입 이미지는 EXIF 방향을 반영한 뒤 긴 변 IMAGE_MAX_PX를 넘을 때만 같은 형식으로 축소한다(HTML·출력 용량 상한). 방향·크기가 그대로면 검증한
  바이트를 그대로 넣는다. 스냅샷의 검증 바이트·content_hash는 어느 경우에도 바뀌지 않는다. 이미지 검증은 전체 픽셀을 디코딩한다.
- 스냅샷은 Document의 Page/Block/content를 깊은 복사한다. 호출자가 원본을 바꿔도 스냅샷·해시·출력은 그대로다.
- 측정 JSON은 load·fonts.ready 이후에 쓰이고, 모양·page_id 집합을 스냅샷과 대조한 뒤에만 쓴다(문서 값으로 위조 불가).
- AI를 호출하지 않는다. 실패는 RenderError(code)로 알리고 부분 파일을 남기지 않는다.
"""
from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import re
import secrets
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from app.config import Settings
from app.models import Document, Page
from app.services import layout_checks

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_DIR = ROOT / "templates" / "export"
TEMPLATE_FILE = TEMPLATE_DIR / "company_intro_v0.html"
FONT_DIR = ROOT / "templates" / "fonts"
FONT_FILES = {"regular": FONT_DIR / "Pretendard-Regular.ttf", "bold": FONT_DIR / "Pretendard-Bold.ttf"}

FORMATS = ("pdf", "docx")
CHECK_KEYS = ("overflow", "broken_image", "placeholder_remaining")
# 형식별 required 검사. 사용자 결정(2026-09-26): DOCX overflow도 required — 엔진이 없으면 not_checked·layout_ok=False.
REQUIRED_CHECKS: dict[str, frozenset[str]] = {"pdf": frozenset(CHECK_KEYS), "docx": frozenset(CHECK_KEYS)}

# 템플릿 v0 배치 상수(형식 공통). 바꾸면 TEMPLATE_VERSION을 올린다(가드 테스트가 아래 값도 해시에 넣는다).
PAGE_W_MM, PAGE_H_MM = 210, 297
IMAGE_MAX_H_MM = 120          # contain: 폭 180mm·높이 120mm 안에 비율 유지
IMAGE_CROP_H_MM = 100         # crop(PDF): 180mm × 100mm 상자에 object-fit: cover. DOCX는 contain으로 대체(제한 기록)
HEADING_PT = {1: 20, 2: 14, 3: 12}
LABEL_PT, CAPTION_PT = 8, 9
IMAGE_MAX_PX = 1600           # 삽입용 이미지의 긴 변 상한(px). 넘으면 같은 형식으로 축소 사본을 넣는다(스냅샷 바이트는 그대로 검증·보관)
DOCX_BOX_SPACE_PT = 18
DOCX_COLORS = {"text": "111111", "label": "8A8A8A", "caption": "555555", "box": "666666", "broken": "A93226"}
DOCX_LAYOUT_CONSTANTS = {"page_mm": [PAGE_W_MM, PAGE_H_MM], "image_max_h_mm": IMAGE_MAX_H_MM, "image_max_px": IMAGE_MAX_PX,
                         "heading_pt": HEADING_PT, "label_pt": LABEL_PT, "caption_pt": CAPTION_PT, "page_break": "per_logical_page",
                         "box": "table_grid_1x1", "box_space_pt": DOCX_BOX_SPACE_PT, "colors": DOCX_COLORS, "crop": "contain_fallback",
                         "exif_orientation": "apply_before_embed", "image_decode": "full_pixels"}
# 배치에 영향을 주는 브라우저 인자(창 크기·가상 시간). 바꾸면 TEMPLATE_VERSION을 올린다(지문 포함). 샌드박스·프로필 등 환경 인자는 제외.
PDF_RENDER_CONSTANTS = {"window_size": "1000,1400", "virtual_time_budget_ms": 10000, "measure": "after_load_and_fonts_ready"}

_BROWSER_CANDIDATES_WIN = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    r"%LOCALAPPDATA%\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
]
_BROWSER_CANDIDATES_POSIX = ["google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "microsoft-edge",
                             "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
                             "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge"]
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ufffe\uffff\ud800-\udfff]")  # C0 제어·DEL·XML 비호환(U+FFFE/FFFF·서로게이트)


class RenderError(Exception):
    """파일을 만들 수 없을 때. BE-08은 이를 EXPORT_FAILED 등으로 바꾼다. code: unsupported_format / browser_not_found /
    render_timeout / render_failed / save_failed / template_missing"""

    def __init__(self, code: str, message: str, details: dict[str, Any] | None = None) -> None:
        super().__init__(f"[{code}] {message}")
        self.code = code
        self.message = message
        self.details = details or {}


# ---------------- 스냅샷 ----------------

@dataclass(frozen=True)
class SnapshotAsset:
    """서버가 검증한 이미지 1개. ok=True면 data는 content_hash와 일치하고 Pillow로 열린 바이트다."""
    asset_id: str
    content_hash: str            # assets.content_hash(승인 manifest에 쓰는 값). 행이 없으면 ""
    mime_type: str
    width: int
    height: int
    data: bytes | None
    ok: bool
    reason: str | None = None    # not_found / not_accessible / not_ready / missing / hash_mismatch / decode_failed


@dataclass(frozen=True)
class RenderSnapshot:
    document_id: str
    document_revision: int
    input_revision: int
    title: str
    target_pages: int
    pages: list[Page]
    assets: dict[str, SnapshotAsset]
    content_hash: str            # {"title","pages"} 정규화 JSON sha256 — 정보용(계약 필드 아님)

    @property
    def manifest_items(self) -> list[tuple[str, str]]:
        """image 블록 순서대로 (asset_id, content_hash). layout_checks.asset_manifest_hash와 같은 입력."""
        items = []
        for page in self.pages:
            for block in page.blocks:
                if block.type == "image":
                    aid = block.content.get("asset_id")
                    asset = self.assets.get(aid)
                    items.append((aid, asset.content_hash if asset else ""))
        return items

    @property
    def asset_manifest_hash(self) -> str:
        return layout_checks.manifest_hash(self.manifest_items)


def _document_content_hash(title: str, pages: list[Page]) -> str:
    payload = {"title": title, "pages": [p.model_dump() for p in pages]}
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()


def asset_from_bytes(asset_id: str, data: bytes | None, *, mime_type: str = "image/png", content_hash: str | None = None,
                     reason: str | None = None) -> SnapshotAsset:
    """바이트를 검증해 SnapshotAsset을 만든다. content_hash를 주면 sha256 일치를 확인한다(불일치 → hash_mismatch).
    data가 None이면 reason(기본 missing)의 실패 asset. 테스트·실험·build_snapshot이 공용으로 쓴다."""
    if data is None:
        return SnapshotAsset(asset_id, content_hash or "", mime_type, 0, 0, None, False, reason or "missing")
    digest = hashlib.sha256(data).hexdigest()
    if content_hash is not None and digest != content_hash:
        return SnapshotAsset(asset_id, content_hash, mime_type, 0, 0, None, False, "hash_mismatch")
    try:
        from PIL import Image

        with Image.open(io.BytesIO(data)) as probe:
            probe.verify()
        with Image.open(io.BytesIO(data)) as img:
            width, height, fmt = img.width, img.height, img.format
            img.load()   # 전체 픽셀 디코딩. 헤더는 정상이지만 본문이 잘린 파일은 여기서 실패한다(LOAD_TRUNCATED_IMAGES 기본 False)
    except Exception:
        return SnapshotAsset(asset_id, content_hash or digest, mime_type, 0, 0, None, False, "decode_failed")
    mime = {"PNG": "image/png", "JPEG": "image/jpeg"}.get(fmt or "")
    if mime is None:  # Pillow는 열지만 브라우저/python-docx가 모두 지원한다고 보장할 수 없는 형식(WebP 등)
        return SnapshotAsset(asset_id, content_hash or digest, mime_type, width, height, None, False, "unsupported_format")
    return SnapshotAsset(asset_id, content_hash or digest, mime, width, height, data, True, None)


def snapshot_from_document(document: Document, assets: dict[str, SnapshotAsset]) -> RenderSnapshot:
    """DB 없이 스냅샷을 만든다(테스트·실험용). 문서가 참조하는데 assets에 없는 asset_id는 not_found로 채운다."""
    assets = dict(assets)
    for aid in layout_checks.image_asset_ids(document):
        if aid not in assets:
            assets[aid] = SnapshotAsset(aid, "", "image/png", 0, 0, None, False, "not_found")
    return RenderSnapshot(document.document_id, document.document_revision, document.input_revision, document.title,
                          document.target_pages, _copy_pages(document.pages), assets, _document_content_hash(document.title, document.pages))


def _copy_pages(pages: list[Page]) -> list[Page]:
    """Page/Block/content dict까지 깊은 복사. 호출자가 원본 Document를 바꿔도 스냅샷·해시·출력이 유지된다."""
    return [page.model_copy(deep=True) for page in pages]


def build_snapshot(conn: sqlite3.Connection, settings: Settings, session_id: str, document: Document) -> RenderSnapshot:
    """이 모듈에서 유일하게 DB·파일을 읽는 함수. image 블록의 asset을 확인하고 바이트를 고정한다.

    - content_hash는 승인 검사(layout_checks.asset_manifest_hash)와 같은 조회(assets.content_hash, 행 없으면 "")로 가져온다.
    - 접근 규칙은 assets.get_ready와 같다: 이 세션의 session asset 또는 registered asset, deleted_at 없음, status=ready.
    - 파일은 settings.private_runs_dir/stored_path에서 읽고 sha256이 content_hash와 같아야 한다(다른 파일로 바꿔치기 → hash_mismatch).
    """
    from app.services.sources import resolve_path

    assets: dict[str, SnapshotAsset] = {}
    for aid in layout_checks.image_asset_ids(document):
        if aid in assets:
            continue
        row = conn.execute("SELECT * FROM assets WHERE asset_id=?", (aid,)).fetchone()
        if row is None:
            assets[aid] = SnapshotAsset(aid, "", "image/png", 0, 0, None, False, "not_found")
            continue
        chash, mime = row["content_hash"], row["mime_type"]
        accessible = row["deleted_at"] is None and (row["scope"] == "registered" or (row["scope"] == "session" and row["session_id"] == session_id))
        if not accessible:
            assets[aid] = SnapshotAsset(aid, chash, mime, 0, 0, None, False, "not_accessible")
            continue
        if row["status"] != "ready":
            assets[aid] = SnapshotAsset(aid, chash, mime, 0, 0, None, False, "not_ready")
            continue
        path = resolve_path(settings, row["stored_path"])
        try:
            data = path.read_bytes()
        except OSError:
            assets[aid] = SnapshotAsset(aid, chash, mime, 0, 0, None, False, "missing")
            continue
        assets[aid] = asset_from_bytes(aid, data, mime_type=mime, content_hash=chash)
    return RenderSnapshot(document.document_id, document.document_revision, document.input_revision, document.title,
                          document.target_pages, _copy_pages(document.pages), assets, _document_content_hash(document.title, document.pages))


# ---------------- 결과 ----------------

@dataclass
class Finding:
    kind: Literal["overflow", "broken_image", "placeholder_remaining"]
    page_id: str
    block_id: str | None
    message: str
    details: dict[str, Any] = field(default_factory=dict)


@dataclass
class LayoutCheckRecord:
    """배치 검사 1종의 결과(계약 확인 ㉜ 제안). required 검사가 not_checked면 layout_ok가 될 수 없다."""
    check_key: str
    required: bool
    result: Literal["ok", "finding", "not_checked"]
    block_ids: list[str] = field(default_factory=list)
    page_ids: list[str] = field(default_factory=list)
    reason: str | None = None


@dataclass
class RenderResult:
    format: str
    file_path: Path
    actual_pages: int | None
    checks: list[LayoutCheckRecord]
    findings: list[Finding]
    layout_ok: bool
    template_version: str
    render_options_hash: str
    asset_manifest_hash: str
    renderer: str
    elapsed_ms: int
    details: dict[str, Any] = field(default_factory=dict)

    @property
    def not_checked(self) -> list[str]:
        return [c.check_key for c in self.checks if c.result == "not_checked"]

    def to_dict(self) -> dict[str, Any]:
        return {"format": self.format, "file_path": str(self.file_path), "actual_pages": self.actual_pages,
                "layout_ok": self.layout_ok, "template_version": self.template_version,
                "render_options_hash": self.render_options_hash, "asset_manifest_hash": self.asset_manifest_hash,
                "renderer": self.renderer, "elapsed_ms": self.elapsed_ms,
                "checks": [c.__dict__ for c in self.checks], "findings": [f.__dict__ for f in self.findings],
                "details": self.details}


def _record(check_key: str, fmt: str, findings: list[Finding], *, not_checked_reason: str | None = None) -> LayoutCheckRecord:
    required = check_key in REQUIRED_CHECKS[fmt]
    if not_checked_reason is not None:
        return LayoutCheckRecord(check_key, required, "not_checked", reason=not_checked_reason)
    mine = [f for f in findings if f.kind == check_key]
    return LayoutCheckRecord(check_key, required, "finding" if mine else "ok",
                             block_ids=[f.block_id for f in mine if f.block_id], page_ids=sorted({f.page_id for f in mine}))


def _layout_ok(checks: list[LayoutCheckRecord]) -> bool:
    return all(c.result == "ok" for c in checks if c.required)


def _snapshot_findings(snapshot: RenderSnapshot) -> list[Finding]:
    """형식과 무관한 검사: 남은 사진 자리, 깨진/없는 이미지."""
    out: list[Finding] = []
    for page in snapshot.pages:
        for block in page.blocks:
            if block.type == "image_placeholder":
                out.append(Finding("placeholder_remaining", page.page_id, block.block_id, "사진 자리가 남아 있습니다. 실제 사진을 넣거나 자리를 제거해 주세요.",
                                   {"description": block.content.get("description", "")}))
            elif block.type == "image":
                aid = block.content.get("asset_id")
                asset = snapshot.assets.get(aid)
                if asset is None or not asset.ok:
                    out.append(Finding("broken_image", page.page_id, block.block_id, "이미지를 열 수 없어 출력에 넣지 못했습니다.",
                                       {"asset_id": aid, "reason": (asset.reason if asset else "not_found")}))
    return out


# ---------------- 블록 뷰 모델(HTML·DOCX 공용) ----------------

EXIF_ORIENTATION_TAG = 0x0112


def embed_bytes(asset: SnapshotAsset) -> tuple[bytes, int, int]:
    """삽입할 (바이트, 폭, 높이). 스냅샷(content_hash·원본 바이트)은 바뀌지 않는다.

    - EXIF Orientation이 1이 아니면 먼저 픽셀을 실제 방향으로 돌린다(PDF·DOCX 모두 같은 방향으로 보이게). 돌린 사본에는 Orientation 태그를 남기지 않는다.
    - 그 뒤 긴 변이 IMAGE_MAX_PX를 넘으면 같은 형식으로 축소한다.
    - 방향 그대로이고 상한 이하면 검증한 바이트를 그대로 돌려준다.
    - 같은 입력·같은 Pillow 버전·같은 상수면 출력 바이트가 같다(PNG optimize·JPEG quality 고정, 시각 메타데이터 없음).
    """
    assert asset.data is not None
    from PIL import Image, ImageOps

    with Image.open(io.BytesIO(asset.data)) as img:
        fmt = img.format or "PNG"
        orientation = img.getexif().get(EXIF_ORIENTATION_TAG, 1)
        needs_transpose = orientation not in (None, 1)
        oriented = ImageOps.exif_transpose(img) if needs_transpose else img
        longest = max(oriented.width, oriented.height)
        if not needs_transpose and longest <= IMAGE_MAX_PX:
            return asset.data, oriented.width, oriented.height
        work = oriented.copy()
    if longest > IMAGE_MAX_PX:
        work.thumbnail((IMAGE_MAX_PX, IMAGE_MAX_PX))
    if fmt == "JPEG" and work.mode not in ("RGB", "L"):
        work = work.convert("RGB")
    buf = io.BytesIO()
    if fmt == "JPEG":
        work.save(buf, format="JPEG", quality=88, optimize=True)   # exif를 넘기지 않으므로 Orientation 태그가 남지 않는다
    else:
        work.save(buf, format="PNG", optimize=True)
    return buf.getvalue(), work.width, work.height


def _clean_text(value: Any) -> str:
    return _CONTROL_CHARS.sub("", str(value if value is not None else ""))


def _view_blocks(snapshot: RenderSnapshot, page: Page) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for block in page.blocks:
        c = block.content
        if block.type == "heading":
            level = c.get("level")
            level = level if level in (1, 2, 3) else 2
            out.append({"type": "heading", "block_id": block.block_id, "level": level, "text": _clean_text(c.get("text"))})
        elif block.type == "paragraph":
            out.append({"type": "paragraph", "block_id": block.block_id, "text": _clean_text(c.get("text"))})
        elif block.type == "list":
            items = c.get("items") if isinstance(c.get("items"), list) else []
            out.append({"type": "list", "block_id": block.block_id, "list_items": [_clean_text(i) for i in items]})
        elif block.type == "image":
            aid = c.get("asset_id")
            asset = snapshot.assets.get(aid)
            caption, alt = _clean_text(c.get("caption")), _clean_text(c.get("alt"))
            if asset is None or not asset.ok or asset.data is None:
                out.append({"type": "broken_image", "block_id": block.block_id, "asset_id": _clean_text(aid),
                            "reason": (asset.reason if asset else "not_found"), "caption": caption})
            else:
                data, width, height = embed_bytes(asset)
                out.append({"type": "image", "block_id": block.block_id, "asset": asset, "fit": "crop" if c.get("fit") == "crop" else "contain",
                            "caption": caption, "alt": alt, "mime_type": asset.mime_type, "embed": data, "width": width, "height": height,
                            "data_b64": base64.b64encode(data).decode("ascii")})
        elif block.type == "image_placeholder":
            out.append({"type": "image_placeholder", "block_id": block.block_id, "description": _clean_text(c.get("description"))})
    return out


# ---------------- HTML(PDF) ----------------

def _font_b64(path: Path) -> str:
    try:
        return base64.b64encode(path.read_bytes()).decode("ascii")
    except OSError as exc:
        raise RenderError("template_missing", f"동봉 폰트를 읽을 수 없습니다: {path.name}") from exc


def build_html(snapshot: RenderSnapshot) -> str:
    """스냅샷 → 단일 HTML(폰트·이미지 data URI, 외부 리소스 없음). 실험 스크립트도 이 함수를 쓴다."""
    from jinja2 import Environment, FileSystemLoader, select_autoescape

    if not TEMPLATE_FILE.is_file():
        raise RenderError("template_missing", f"템플릿이 없습니다: {TEMPLATE_FILE.name}")
    env = Environment(loader=FileSystemLoader(str(TEMPLATE_DIR)), autoescape=select_autoescape(default=True, default_for_string=True))
    opts = layout_checks.DEFAULT_RENDER_OPTIONS
    margin = opts["margin_mm"]
    template = env.get_template(TEMPLATE_FILE.name)
    return template.render(
        title=_clean_text(snapshot.title), nonce=secrets.token_urlsafe(16),
        font_family=opts["font_family"], font_regular_b64=_font_b64(FONT_FILES["regular"]), font_bold_b64=_font_b64(FONT_FILES["bold"]),
        page_size=opts["page_size"], margin_mm=margin, base_font_pt=opts["base_font_pt"],
        content_w_mm=PAGE_W_MM - 2 * margin, content_h_mm=PAGE_H_MM - 2 * margin,
        image_max_h_mm=IMAGE_MAX_H_MM, image_crop_h_mm=IMAGE_CROP_H_MM,
        pages=[{"page_id": p.page_id, "title": _clean_text(p.title), "blocks": _view_blocks(snapshot, p)} for p in snapshot.pages],
    )


def find_browser(settings: Settings | None = None) -> Path | None:
    """EXPORT_BROWSER_PATH → Chrome → Edge 순서로 실행 파일을 찾는다. 없으면 None."""
    if settings is not None and settings.export_browser_path:
        p = Path(os.path.expandvars(settings.export_browser_path))
        return p if p.is_file() else None
    if sys.platform == "win32":
        for raw in _BROWSER_CANDIDATES_WIN:
            p = Path(os.path.expandvars(raw))
            if p.is_file():
                return p
        return None
    for name in _BROWSER_CANDIDATES_POSIX:
        if name.startswith("/"):
            if Path(name).is_file():
                return Path(name)
        elif (found := shutil.which(name)):
            return Path(found)
    return None


def browser_version(path: Path) -> str:
    """'chrome/153.0.8010.54' 형태. Windows는 파일 버전 정보, 그 외는 --version."""
    kind = "edge" if "edge" in path.name.lower() else ("chrome" if "chrome" in path.name.lower() else path.stem.lower())
    version = "unknown"
    try:
        if sys.platform == "win32":
            version = _win_file_version(path) or version
        else:
            cp = subprocess.run([str(path), "--version"], capture_output=True, text=True, timeout=15)
            m = re.search(r"(\d+\.\d+\.\d+\.\d+)", cp.stdout + cp.stderr)
            version = m.group(1) if m else version
    except Exception:
        pass
    return f"{kind}/{version}"


def _win_file_version(path: Path) -> str | None:
    import ctypes

    size = ctypes.windll.version.GetFileVersionInfoSizeW(str(path), None)
    if not size:
        return None
    buf = ctypes.create_string_buffer(size)
    if not ctypes.windll.version.GetFileVersionInfoW(str(path), 0, size, buf):
        return None
    ptr = ctypes.c_void_p()
    length = ctypes.c_uint()
    if not ctypes.windll.version.VerQueryValueW(buf, "\\", ctypes.byref(ptr), ctypes.byref(length)):
        return None
    fixed = ctypes.cast(ptr, ctypes.POINTER(ctypes.c_uint32 * 13)).contents
    ms, ls = fixed[2], fixed[3]
    return f"{ms >> 16}.{ms & 0xFFFF}.{ls >> 16}.{ls & 0xFFFF}"


def _browser_args(browser: Path, profile_dir: Path) -> list[str]:
    args = [str(browser), "--headless=new", "--disable-gpu", "--no-first-run", "--no-default-browser-check",
            "--disable-extensions", "--disable-sync", "--disable-background-networking", "--disable-component-update",
            f"--user-data-dir={profile_dir}", f"--window-size={PDF_RENDER_CONSTANTS['window_size']}", "--hide-scrollbars",
            f"--virtual-time-budget={PDF_RENDER_CONSTANTS['virtual_time_budget_ms']}"]
    # 샌드박스를 끄는 인자(--no-sandbox)는 넣지 않는다. Chromium은 root로 실행되면 샌드박스 때문에 시작을 거부하므로
    # 서버는 비root 사용자로 실행해야 한다(README·task_backend.md 6.8-6).
    return args


def _kill_tree(proc: subprocess.Popen) -> None:
    try:
        if sys.platform == "win32":
            subprocess.run(["taskkill", "/T", "/F", "/PID", str(proc.pid)], capture_output=True, timeout=15)
        else:
            import signal

            os.killpg(os.getpgid(proc.pid), signal.SIGKILL)
    except Exception:
        pass
    try:
        proc.kill()
    except Exception:
        pass


def _run(cmd: list[str], timeout: int, what: str) -> subprocess.CompletedProcess:
    """브라우저를 실행하고 stdout/stderr를 모은다. 시간 초과 시 자식 프로세스까지 종료한다."""
    try:
        proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                **({} if sys.platform == "win32" else {"start_new_session": True}))
    except OSError as exc:
        raise RenderError("render_failed", f"브라우저를 실행할 수 없습니다: {what}") from exc
    try:
        out, err = proc.communicate(timeout=timeout)
    except subprocess.TimeoutExpired as exc:
        _kill_tree(proc)
        try:
            proc.communicate(timeout=15)
        except Exception:
            pass
        raise RenderError("render_timeout", f"브라우저 {what} 단계가 {timeout}초 안에 끝나지 않았습니다.") from exc
    return subprocess.CompletedProcess(cmd, proc.returncode, out, err)


def print_pdf_and_measure(browser: Path, html_path: Path, pdf_path: Path, profile_dir: Path, timeout: int) -> dict[str, Any] | None:
    """브라우저 1회 실행으로 PDF(--print-to-pdf)와 측정 DOM(--dump-dom)을 함께 얻는다.

    PDF가 없으면 RenderError(render_failed). 측정 JSON(#be07-measure)을 읽지 못하면 None을 돌려주고 호출자가 overflow를 not_checked로 남긴다.
    """
    cp = _run(_browser_args(browser, profile_dir) + ["--dump-dom", "--no-pdf-header-footer", f"--print-to-pdf={pdf_path}", html_path.as_uri()],
              timeout, "print-to-pdf")
    if cp.returncode != 0 or not pdf_path.is_file() or pdf_path.stat().st_size == 0:
        raise RenderError("render_failed", "PDF가 만들어지지 않았습니다.", {"returncode": cp.returncode})
    dom = cp.stdout.decode("utf-8", "replace")
    m = re.search(r'<script type="application/json" id="be07-measure">(.*?)</script>', dom, re.S)
    if not m:
        return None
    try:
        # script 요소 본문은 직렬화 때 엔티티 변환이 없으므로 그대로 JSON이다(unescape하면 ID 안의 &…; 가 바뀌어 위조·손상 가능).
        return json.loads(m.group(1))
    except ValueError:
        return None


def _valid_measure(measure: Any, snapshot: RenderSnapshot) -> bool:
    """측정 JSON의 모양과 page_id 집합이 스냅샷과 일치하는지. 문서 값(ID)이 JSON 구조를 흉내 내도 통과하지 못한다."""
    if not isinstance(measure, dict) or not isinstance(measure.get("pages"), list) or measure.get("fonts_ready") is not True:
        return False
    expected = [p.page_id for p in snapshot.pages]
    got = []
    block_ids = {b.block_id for p in snapshot.pages for b in p.blocks}
    for p in measure["pages"]:
        if not isinstance(p, dict) or not isinstance(p.get("page_id"), str) or not isinstance(p.get("overflow"), bool):
            return False
        if not isinstance(p.get("height_px"), (int, float)) or not isinstance(p.get("excess_px"), (int, float)):
            return False
        first = p.get("first_overflow_block_id")
        if first is not None and (not isinstance(first, str) or first not in block_ids):
            return False
        got.append(p["page_id"])
    return got == expected


def pdf_info(pdf_path: Path) -> dict[str, Any]:
    """pypdf로 쪽수·1쪽 텍스트·폰트 이름을 읽는다(구조 검사)."""
    from pypdf import PdfReader

    reader = PdfReader(str(pdf_path))
    fonts: set[str] = set()
    for page in reader.pages:
        res = page.get("/Resources") or {}
        for _, f in (res.get("/Font") or {}).items():
            try:
                fonts.add(str(f.get_object().get("/BaseFont")))
            except Exception:
                pass
    return {"pages": len(reader.pages), "fonts": sorted(fonts),
            "first_page_text": (reader.pages[0].extract_text() or "") if reader.pages else ""}


def _overflow_findings(measure: dict[str, Any]) -> list[Finding]:
    out = []
    for p in measure.get("pages", []):
        if p.get("overflow"):
            excess_mm = round(float(p.get("excess_px", 0)) * 25.4 / 96, 1)
            out.append(Finding("overflow", p["page_id"], p.get("first_overflow_block_id"),
                               f"이 페이지 내용이 A4 한 쪽 본문 높이를 {excess_mm}mm 넘칩니다.",
                               {"excess_mm": excess_mm, "height_mm": round(float(p.get("height_px", 0)) * 25.4 / 96, 1)}))
    return out


def _render_pdf(snapshot: RenderSnapshot, out_dir: Path, settings: Settings | None, timeout: int) -> tuple[Path, int | None, list[Finding], list[LayoutCheckRecord], str, dict[str, Any]]:
    browser = find_browser(settings)
    if browser is None:
        raise RenderError("browser_not_found", "PDF 생성에 필요한 Chromium 계열 브라우저(Chrome/Edge)를 찾지 못했습니다. EXPORT_BROWSER_PATH를 설정해 주세요.")
    renderer = browser_version(browser)
    final = out_dir / f"{snapshot.document_id}_rev{snapshot.document_revision}.pdf"
    html = build_html(snapshot)
    findings = _snapshot_findings(snapshot)
    details: dict[str, Any] = {"browser_path": str(browser)}
    with tempfile.TemporaryDirectory(prefix=".render_", dir=out_dir) as tmp:
        tmp_dir = Path(tmp)
        html_path = tmp_dir / "page.html"
        html_path.write_text(html, encoding="utf-8")
        pdf_tmp = tmp_dir / "out.pdf"
        measure = print_pdf_and_measure(browser, html_path, pdf_tmp, tmp_dir / "profile", timeout)
        overflow_reason: str | None = None
        if measure is None:
            overflow_reason = "measure_failed"
        elif not _valid_measure(measure, snapshot):
            overflow_reason = "measure_invalid"
        else:
            findings += _overflow_findings(measure)
            details["measure"] = measure
        try:
            info = pdf_info(pdf_tmp)   # 발행 전에 구조 검사. 읽지 못하는 파일은 내보내지 않는다
        except Exception as exc:  # noqa: BLE001
            raise RenderError("render_failed", "만들어진 PDF를 읽을 수 없습니다.", {"error": type(exc).__name__}) from exc
        try:
            os.replace(pdf_tmp, final)
        except OSError as exc:
            raise RenderError("save_failed", f"파일 저장에 실패했습니다: {final.name}") from exc
    details["fonts"] = info["fonts"]
    checks = [_record("overflow", "pdf", findings, not_checked_reason=overflow_reason),
              _record("broken_image", "pdf", findings), _record("placeholder_remaining", "pdf", findings)]
    return final, info["pages"], findings, checks, renderer, details


# ---------------- DOCX ----------------

def _docx_set_font(style, family: str, size_pt: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    style.font.name = family
    if size_pt is not None:
        style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), family)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


def _docx_box(doc, title: str, body: str, broken: bool) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(DOCX_COLORS["broken"] if broken else DOCX_COLORS["box"])
    para2 = cell.add_paragraph(body)
    para2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(DOCX_BOX_SPACE_PT)
        p.paragraph_format.space_after = Pt(DOCX_BOX_SPACE_PT)
    doc.add_paragraph()


def _fit_mm(width_px: int, height_px: int, max_w_mm: float, max_h_mm: float) -> tuple[float, float]:
    scale = min(max_w_mm / width_px, max_h_mm / height_px)
    return width_px * scale, height_px * scale


def _render_docx(snapshot: RenderSnapshot, out_dir: Path) -> tuple[Path, int | None, list[Finding], list[LayoutCheckRecord], str, dict[str, Any]]:
    opts = layout_checks.DEFAULT_RENDER_OPTIONS
    family, margin = opts["font_family"], opts["margin_mm"]
    content_w = PAGE_W_MM - 2 * margin
    final = out_dir / f"{snapshot.document_id}_rev{snapshot.document_revision}.docx"
    findings = _snapshot_findings(snapshot)

    try:
        doc = _build_docx(snapshot, family, opts["base_font_pt"], content_w)
    except RenderError:
        raise
    except Exception as exc:  # noqa: BLE001  (lxml ValueError, python-docx 이미지 오류 등) → 500이 아니라 RenderError
        raise RenderError("render_failed", "DOCX 본문을 만들지 못했습니다.", {"error": type(exc).__name__}) from exc

    out_dir.mkdir(parents=True, exist_ok=True)
    tmp = out_dir / f".{final.stem}.{secrets.token_hex(6)}.tmp"   # 같은 문서를 동시에 만들어도 서로 다른 임시 파일
    try:
        doc.save(str(tmp))
        if not tmp.is_file() or tmp.stat().st_size == 0:
            raise RenderError("save_failed", f"파일이 생성되지 않았습니다: {final.name}")
        try:
            details = docx_info(tmp)   # 발행 전에 다시 열어 구조 확인
        except Exception as exc:  # noqa: BLE001
            raise RenderError("render_failed", "만들어진 DOCX를 다시 열 수 없습니다.", {"error": type(exc).__name__}) from exc
        os.replace(tmp, final)
    except OSError as exc:
        raise RenderError("save_failed", f"파일 저장에 실패했습니다: {final.name}") from exc
    finally:
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass
    details["font_embedding"] = "not_supported_in_this_implementation"
    checks = [_record("overflow", "docx", findings, not_checked_reason="docx_no_layout_engine"),
              _record("broken_image", "docx", findings), _record("placeholder_remaining", "docx", findings)]
    return final, None, findings, checks, "python-docx/" + _module_version("docx"), details


def _build_docx(snapshot: RenderSnapshot, family: str, base_pt: float, content_w: float):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor

    doc = docx.Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(PAGE_W_MM), Mm(PAGE_H_MM)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Mm(layout_checks.DEFAULT_RENDER_OPTIONS["margin_mm"])
    _docx_set_font(doc.styles["Normal"], family, base_pt, color=DOCX_COLORS["text"])
    for level, size in HEADING_PT.items():
        _docx_set_font(doc.styles[f"Heading {level}"], family, size, bold=True, color=DOCX_COLORS["text"])
    _docx_set_font(doc.styles["Title"], family, HEADING_PT[1], bold=True, color=DOCX_COLORS["text"])
    _docx_set_font(doc.styles["Caption"], family, CAPTION_PT, bold=False, color=DOCX_COLORS["caption"])
    _docx_set_font(doc.styles["List Bullet"], family, base_pt)
    doc.core_properties.title = _clean_text(snapshot.title)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""

    for index, page in enumerate(snapshot.pages, start=1):
        if index > 1:
            doc.add_page_break()
        label = doc.add_paragraph()
        label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = label.add_run(f"{index} / {_clean_text(page.title)}")
        run.font.size = Pt(LABEL_PT)
        run.font.color.rgb = RGBColor.from_string(DOCX_COLORS["label"])
        for b in _view_blocks(snapshot, page):
            if b["type"] == "heading":
                doc.add_heading(b["text"], level=b["level"])
            elif b["type"] == "paragraph":
                doc.add_paragraph(b["text"])
            elif b["type"] == "list":
                for item in b["list_items"]:
                    doc.add_paragraph(item, style="List Bullet")
            elif b["type"] == "image":
                w_mm, h_mm = _fit_mm(b["width"], b["height"], content_w, IMAGE_MAX_H_MM)
                doc.add_picture(io.BytesIO(b["embed"]), width=Mm(w_mm), height=Mm(h_mm))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if b["caption"]:
                    cap = doc.add_paragraph(b["caption"], style="Caption")
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif b["type"] == "broken_image":
                _docx_box(doc, "[이미지를 열 수 없음]", f"{b['asset_id']} ({b['reason']})" + (f"\n{b['caption']}" if b["caption"] else ""), True)
            elif b["type"] == "image_placeholder":
                _docx_box(doc, "[사진 자리]", b["description"], False)
    return doc


def docx_info(path: Path) -> dict[str, Any]:
    """python-docx로 다시 열어 구조만 센다(쪽수·넘침 판단에 쓰지 않는다)."""
    import docx

    d = docx.Document(str(path))
    return {"paragraphs": len(d.paragraphs), "inline_shapes": len(d.inline_shapes), "tables": len(d.tables),
            "sections": len(d.sections), "text_chars": sum(len(p.text) for p in d.paragraphs)}


def _module_version(name: str) -> str:
    try:
        from importlib.metadata import version

        return version("python-docx" if name == "docx" else name)
    except Exception:
        return "unknown"


# ---------------- 공개 함수 ----------------

def render(snapshot: RenderSnapshot, fmt: str, out_dir: Path, settings: Settings | None = None) -> RenderResult:
    """스냅샷을 out_dir 아래 실제 파일로 만든다. 실패 시 RenderError(성공 값을 임의로 만들지 않는다).

    settings는 브라우저 경로·시간 제한에만 쓴다(렌더 옵션이 아니다). None이면 자동 탐색·기본 90초.
    """
    if fmt not in FORMATS:
        raise RenderError("unsupported_format", f"지원하지 않는 형식입니다: {fmt}")
    out_dir = Path(out_dir)
    try:
        out_dir.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise RenderError("save_failed", f"저장 폴더를 만들 수 없습니다: {out_dir}") from exc
    timeout = settings.export_render_timeout_s if settings is not None else 90
    started = time.perf_counter()
    if fmt == "pdf":
        path, pages, findings, checks, renderer, details = _render_pdf(snapshot, out_dir, settings, timeout)
    else:
        path, pages, findings, checks, renderer, details = _render_docx(snapshot, out_dir)
    return RenderResult(
        format=fmt, file_path=path, actual_pages=pages, checks=checks, findings=findings, layout_ok=_layout_ok(checks),
        template_version=layout_checks.TEMPLATE_VERSION, render_options_hash=layout_checks.RENDER_OPTIONS_HASH,
        asset_manifest_hash=snapshot.asset_manifest_hash, renderer=renderer,
        elapsed_ms=int((time.perf_counter() - started) * 1000), details=details,
    )


def template_fingerprint() -> str:
    """템플릿·폰트·DOCX 배치 상수의 sha256. tests/test_be07.py가 TEMPLATE_VERSION별 고정값과 비교한다."""
    h = hashlib.sha256()
    for path in (TEMPLATE_FILE, FONT_FILES["regular"], FONT_FILES["bold"]):
        data = path.read_bytes()
        if path.suffix == ".html":
            data = data.replace(b"\r\n", b"\n")   # 체크아웃 줄바꿈(autocrlf)에 흔들리지 않게
        h.update(path.name.encode())
        h.update(data)
    h.update(json.dumps(DOCX_LAYOUT_CONSTANTS, sort_keys=True).encode())
    h.update(json.dumps(PDF_RENDER_CONSTANTS, sort_keys=True).encode())
    h.update(json.dumps({"image_crop_h_mm": IMAGE_CROP_H_MM, "image_max_px": IMAGE_MAX_PX}, sort_keys=True).encode())
    return h.hexdigest()
