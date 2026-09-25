"""DOCX(python-docx) · PPTX(python-pptx). 읽기 전용이며 출력 도구(D-03) 결정과 무관하다.

DOCX locator: 문단 {"paragraph": n} / 표 셀 {"table": t, "row": r, "col": c}  (모두 1부터)
PPTX locator: 텍스트 상자 {"slide": s, "shape": n} / 표 셀 {"slide": s, "shape": n, "row": r, "col": c}
              그룹 안 도형은 {"slide": s, "shape": n, "child": m}
열 수 없는 파일(손상 또는 암호 — 둘 다 zip으로 열리지 않는다)은 failed + FILE_CORRUPT.
글자가 하나도 없으면: 그림이 있으면 partial + IMAGE_ONLY, 없으면 failed + NO_USABLE_TEXT.
"""
from __future__ import annotations

import io
import logging

from app.parsers import FILE_CORRUPT, IMAGE_ONLY, NO_USABLE_TEXT, ParseResult, Segment, failed, warning

logger = logging.getLogger(__name__)

_CANNOT_OPEN = ("파일을 열 수 없습니다(손상되었거나 암호가 걸려 있습니다).",
                "암호를 풀거나 다시 저장한 파일을 올리거나 이 자료를 제외해 주세요.")


def _has_image_rel(part) -> bool:
    try:
        return any(rel.reltype.endswith("/image") for rel in part.rels.values())
    except Exception:
        return False


def _finish(segments: list[Segment], has_images: bool, image_msg: str) -> ParseResult:
    if segments:
        return ParseResult(status="complete", segments=segments, text_available=True)
    if has_images:
        return ParseResult(status="partial", text_available=False,
                           warnings=[warning(IMAGE_ONLY, image_msg, None,
                                             "텍스트가 들어 있는 원본 파일이나 텍스트본을 추가해 주세요.")])
    return failed(NO_USABLE_TEXT, "읽을 수 있는 글자가 없습니다.", "내용이 있는 파일을 사용해 주세요.")


# ---------------- DOCX ----------------

def parse_docx(data: bytes) -> ParseResult:
    try:
        from docx import Document

        doc = Document(io.BytesIO(data))
    except Exception:
        return failed(FILE_CORRUPT, *_CANNOT_OPEN)

    segments: list[Segment] = []
    for no, para in enumerate(doc.paragraphs, start=1):
        text = para.text.strip()
        if text:
            segments.append(Segment({"paragraph": no}, text))
    for t, table in enumerate(doc.tables, start=1):
        seen_tcs: list = []  # 병합 셀은 같은 xml 요소가 반복된다. 참조를 들고 있어야 같은 객체로 비교된다.
        for r, row in enumerate(table.rows, start=1):
            for c, cell in enumerate(row.cells, start=1):
                if any(cell._tc is tc for tc in seen_tcs):
                    continue
                seen_tcs.append(cell._tc)
                text = cell.text.strip()
                if text:
                    segments.append(Segment({"table": t, "row": r, "col": c}, text))
    return _finish(segments, _has_image_rel(doc.part), "이미지만 있음, 글자를 읽지 못함.")


# ---------------- PPTX ----------------

def _pptx_shape_segments(shape, base: dict, segments: list[Segment]) -> None:
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = shape.text_frame.text.strip()
        if text:
            segments.append(Segment(dict(base), text))
    if getattr(shape, "has_table", False) and shape.has_table:
        for r, row in enumerate(shape.table.rows, start=1):
            for c, cell in enumerate(row.cells, start=1):
                text = cell.text.strip()
                if text:
                    segments.append(Segment({**base, "row": r, "col": c}, text))
    if getattr(shape, "shape_type", None) is not None and hasattr(shape, "shapes"):
        for m, child in enumerate(shape.shapes, start=1):
            _pptx_shape_segments(child, {**base, "child": m}, segments)


def parse_pptx(data: bytes) -> ParseResult:
    try:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(data))
        slides = list(prs.slides)
    except Exception:
        return failed(FILE_CORRUPT, *_CANNOT_OPEN)

    segments: list[Segment] = []
    image_only_slides: list[int] = []
    any_images = False
    for s, slide in enumerate(slides, start=1):
        before = len(segments)
        for n, shape in enumerate(slide.shapes, start=1):
            _pptx_shape_segments(shape, {"slide": s, "shape": n}, segments)
        slide_has_image = _has_image_rel(slide.part)
        any_images = any_images or slide_has_image
        if len(segments) == before and slide_has_image:
            image_only_slides.append(s)

    if not segments:
        return _finish(segments, any_images, "이미지만 있음, 글자를 읽지 못함.")
    warnings = []
    if image_only_slides:
        warnings.append(warning(IMAGE_ONLY, f"{len(image_only_slides)}개 슬라이드는 이미지만 있어 글자를 읽지 못했습니다.",
                                {"slides": image_only_slides}, "해당 슬라이드 내용이 필요하면 텍스트본을 추가해 주세요."))
    return ParseResult(status="partial" if image_only_slides else "complete",
                       segments=segments, warnings=warnings, text_available=True)
